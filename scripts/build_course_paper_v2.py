from __future__ import annotations

import csv
import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont, ImageOps

from build_course_paper import (
    EXP_FIG_DIR,
    FIG_DIR,
    GITHUB_URL,
    HANDDRAWN_PDF,
    PAPER_DIR,
    ROOT,
    TABLE_DIR,
    WORKSPACE,
    add_caption,
    add_figure,
    add_para,
    add_table,
    add_title,
    arrow,
    build_table_data,
    draw_box,
    draw_wrapped,
    fmt2,
    fmt4,
    fmt_params,
    generate_figures,
    main_findings,
    markdown_table,
    pil_font,
    read_csv,
    set_run_font,
    text_size,
)


TITLE = "基于 U-KAN 的医学图像分割算法复现、实验诊断与注意力增强研究"
MD_OUT = PAPER_DIR / "course_paper_draft_v2.md"
DOCX_OUT = PAPER_DIR / "基于U-KAN的医学图像分割算法复现、实验诊断与注意力增强研究_重写版.docx"


def add_heading(doc: Document, text: str, level: int) -> None:
    """Add a visually consistent heading that also uses Word heading styles."""
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 14, True, east_asia="黑体")


KAN_PAPER_DIR = WORKSPACE / "课程论文要求" / "kan相关论文"
OFFICIAL_UKAN_FIG = (
    KAN_PAPER_DIR
    / "U-KAN需要特别关注"
    / "U-KAN-main"
    / "U-KAN-main"
    / "assets"
    / "framework-1.jpg"
)
REVIEW_TABLE_FIG = KAN_PAPER_DIR / "A_Review_of_KANs_for_Medical_Image_Analysis_17.png"


def copy_reference_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    out: dict[str, Path] = {}
    if OFFICIAL_UKAN_FIG.exists():
        target = FIG_DIR / "official_ukan_framework.jpg"
        shutil.copy2(OFFICIAL_UKAN_FIG, target)
        out["official_ukan_framework"] = target
    if REVIEW_TABLE_FIG.exists():
        target = FIG_DIR / "kan_medical_review_table.png"
        shutil.copy2(REVIEW_TABLE_FIG, target)
        out["kan_medical_review_table"] = target
    if HANDDRAWN_PDF.exists():
        shutil.copy2(HANDDRAWN_PDF, PAPER_DIR / HANDDRAWN_PDF.name)
    return out


def load_results() -> list[dict[str, str]]:
    return read_csv(TABLE_DIR / "segmentation_results.csv")


def find_row(rows: list[dict[str, str]], name: str) -> dict[str, str]:
    for row in rows:
        if row["name"] == name:
            return row
    raise KeyError(name)


def paste_contain(canvas: Image.Image, src_path: Path, box: tuple[int, int, int, int], bg: str = "#ffffff") -> None:
    x1, y1, x2, y2 = box
    ImageDraw.Draw(canvas).rounded_rectangle(box, radius=18, fill=bg, outline="#cbd5e1", width=2)
    if not src_path.exists():
        return
    src = Image.open(src_path).convert("RGB")
    src.thumbnail((x2 - x1 - 26, y2 - y1 - 26), Image.Resampling.LANCZOS)
    x = x1 + (x2 - x1 - src.width) // 2
    y = y1 + (y2 - y1 - src.height) // 2
    canvas.paste(src, (x, y))


def create_dataset_combo() -> Path:
    out = FIG_DIR / "dataset_samples_combined.png"
    W, H = 1850, 1050
    img = Image.new("RGB", (W, H), "#f8fafc")
    draw = ImageDraw.Draw(img)
    add_title(draw, "BUSI 与 CVC-ClinicDB 数据样本可视化", W, 35)
    title_font = pil_font(32, True)
    draw.text((95, 125), "BUSI 乳腺超声：目标边界模糊、背景噪声强", font=title_font, fill="#0f172a")
    draw.text((95, 560), "CVC-ClinicDB 结肠镜：息肉形态较规则但边界仍有干扰", font=title_font, fill="#0f172a")
    paste_contain(img, EXP_FIG_DIR / "busi_samples.png", (85, 175, 1765, 505))
    paste_contain(img, EXP_FIG_DIR / "cvc_samples.png", (85, 610, 1765, 940))
    draw.text((90, 985), "图源说明：由本文数据整理脚本从本地处理后的公开数据集中抽样生成。", font=pil_font(20), fill="#475569")
    img.save(out, quality=95)
    return out


def create_results_chart() -> Path:
    out = FIG_DIR / "main_results_iou_dice_chart.png"
    rows = load_results()
    datasets = [
        ("BUSI", ["busi_unet_seed2981", "busi_no_kan_seed2981", "busi_ukan_seed2981", "busi_attention_ukan_seed2981"]),
        ("CVC", ["cvc_unet_seed2981", "cvc_no_kan_seed2981", "cvc_ukan_seed2981", "cvc_attention_ukan_seed2981"]),
    ]
    labels = ["U-Net", "no-KAN", "U-KAN", "Att-U-KAN"]
    colors = ["#2563eb", "#f59e0b", "#059669", "#dc2626"]
    W, H = 1900, 1100
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "主实验 IoU / Dice 对比", W, 35)
    axis_font = pil_font(22)
    value_font = pil_font(18)
    panel_w, panel_h = 780, 720
    for pi, (dataset, names) in enumerate(datasets):
        ox = 130 + pi * 900
        oy = 205
        draw.rounded_rectangle((ox - 40, oy - 80, ox + panel_w + 35, oy + panel_h + 80), radius=24, fill="#ffffff", outline="#cbd5e1", width=2)
        draw.text((ox + 265, oy - 55), dataset, font=pil_font(34, True), fill="#0f172a")
        y0 = oy + panel_h
        x0 = ox + 80
        draw.line((x0, oy, x0, y0), fill="#64748b", width=3)
        draw.line((x0, y0, ox + panel_w, y0), fill="#64748b", width=3)
        for tick in [0.5, 0.6, 0.7, 0.8, 0.9]:
            y = y0 - int((tick - 0.45) / 0.5 * panel_h)
            draw.line((x0 - 8, y, ox + panel_w, y), fill="#e2e8f0", width=1)
            draw.text((ox + 15, y - 12), f"{tick:.1f}", font=axis_font, fill="#475569")
        group_w = 145
        for i, name in enumerate(names):
            r = find_row(rows, name)
            iou = float(r["iou"])
            dice = float(r["dice"])
            gx = x0 + 55 + i * group_w
            for j, (metric, val, shade) in enumerate([("IoU", iou, colors[i]), ("Dice", dice, "#94a3b8")]):
                bw = 45
                h = int((val - 0.45) / 0.5 * panel_h)
                x = gx + j * 52
                draw.rounded_rectangle((x, y0 - h, x + bw, y0), radius=8, fill=shade)
                draw.text((x - 3, y0 - h - 25), f"{val:.3f}", font=value_font, fill="#0f172a")
            tw = text_size(draw, labels[i], axis_font)[0]
            draw.text((gx + 25 - tw // 2, y0 + 18), labels[i], font=axis_font, fill=colors[i])
        draw.text((ox + 495, oy + panel_h + 48), "深色=IoU  灰色=Dice", font=pil_font(20), fill="#475569")
    draw.text((90, 1030), "图源说明：作者根据 evaluate.py 输出的验证集指标整理绘制。", font=pil_font(20), fill="#475569")
    img.save(out, quality=95)
    return out


def create_diagnosis_flow() -> Path:
    out = FIG_DIR / "experiment_diagnosis_flow.png"
    W, H = 2450, 1050
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "从异常曲线到补充实验的思考过程", W, 35)
    body = pil_font(25)
    small = pil_font(20)
    boxes = [
        ("现象 A\nBUSI no-KAN loss 偏高\n但验证 IoU 局部跳升", "#fff7ed", "#ea580c"),
        ("疑问 A\n是否只是阈值化、小验证集\n或随机划分造成的偶然峰值？", "#fef3c7", "#d97706"),
        ("补充实验 A\n更换 seed=6142\n重新训练 no-KAN 与 U-KAN", "#ecfdf5", "#059669"),
        ("结论 A\nU-KAN 仍优于 no-KAN\n但 BUSI 对 seed 敏感", "#e0f2fe", "#0284c7"),
        ("现象 B\nCVC best epoch 靠后\n怀疑 100 epoch 未收敛", "#fce7f3", "#db2777"),
        ("补充实验 B\n从 best checkpoint\n继续训练 50 epoch", "#eef2ff", "#4f46e5"),
        ("结论 B\n未超过原始 best\n差距不能简单归因于没跑够", "#f8fafc", "#64748b"),
    ]
    x, y = 80, 295
    prev = None
    for i, (txt, fill, outline) in enumerate(boxes):
        w = 270 if i not in [1, 2, 6] else 320
        draw_box(draw, (x, y, x + w, y + 185), txt, fill, outline, body)
        if prev:
            arrow(draw, prev, (x, y + 92), "#64748b", 5)
        prev = (x + w, y + 92)
        x += w + 28
    draw_box(
        draw,
        (260, 650, 2190, 845),
        "本文不只报告单次最优指标，而是把训练曲线中的异常作为研究问题：先提出可检验假设，再通过随机种子稳定性实验和继续训练实验进行验证，最终给出更谨慎的复现结论。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw.text((90, 960), "图源说明：作者根据本项目实验讨论、训练曲线观察和补充实验流程自行整理绘制。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def create_supplement_chart() -> Path:
    out = FIG_DIR / "supplemental_experiment_chart.png"
    seed_rows = read_csv(TABLE_DIR / "busi_seed_stability.csv")
    cont_rows = read_csv(TABLE_DIR / "cvc_continued_training.csv")
    W, H = 1800, 1000
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "补充实验结果：稳定性与收敛性验证", W, 35)
    font = pil_font(23)
    value_font = pil_font(20)
    # Left panel: seed stability
    draw.rounded_rectangle((90, 160, 870, 840), radius=24, fill="#ffffff", outline="#cbd5e1", width=2)
    draw.text((240, 190), "BUSI seed 稳定性", font=pil_font(32, True), fill="#0f172a")
    x0, y0 = 180, 745
    draw.line((x0, 300, x0, y0), fill="#64748b", width=3)
    draw.line((x0, y0, 805, y0), fill="#64748b", width=3)
    colors = ["#f59e0b", "#059669"]
    for tick in [0.55, 0.60, 0.65, 0.70]:
        y = y0 - int((tick - 0.52) / 0.2 * 445)
        draw.line((x0 - 8, y, 805, y), fill="#e2e8f0", width=1)
        draw.text((120, y - 11), f"{tick:.2f}", font=value_font, fill="#475569")
    for i, row in enumerate(seed_rows):
        label = row["variant"].replace("U-KAN", "UKAN")
        for j, key in enumerate(["seed2981_iou", "seed6142_iou", "mean_iou"]):
            val = float(row[key])
            bx = x0 + 80 + i * 290 + j * 62
            bh = int((val - 0.52) / 0.2 * 445)
            draw.rounded_rectangle((bx, y0 - bh, bx + 48, y0), radius=8, fill=colors[i])
            draw.text((bx - 3, y0 - bh - 25), f"{val:.3f}", font=value_font, fill="#0f172a")
        draw.text((x0 + 58 + i * 290, y0 + 25), label, font=font, fill=colors[i])
    draw.text((325, 805), "每组：seed2981 / seed6142 / mean", font=value_font, fill="#475569")

    # Right panel: continued training
    draw.rounded_rectangle((930, 160, 1710, 840), radius=24, fill="#ffffff", outline="#cbd5e1", width=2)
    draw.text((1110, 190), "CVC U-KAN 继续训练", font=pil_font(32, True), fill="#0f172a")
    x0, y0 = 1040, 745
    draw.line((x0, 300, x0, y0), fill="#64748b", width=3)
    draw.line((x0, y0, 1645, y0), fill="#64748b", width=3)
    for tick in [0.76, 0.78, 0.80]:
        y = y0 - int((tick - 0.75) / 0.07 * 445)
        draw.line((x0 - 8, y, 1645, y), fill="#e2e8f0", width=1)
        draw.text((970, y - 11), f"{tick:.2f}", font=value_font, fill="#475569")
    labels = ["100 epoch\nbest", "继续训练\n+50 epoch"]
    for i, row in enumerate(cont_rows):
        val = float(row["iou"])
        bx = x0 + 140 + i * 240
        bh = int((val - 0.75) / 0.07 * 445)
        draw.rounded_rectangle((bx, y0 - bh, bx + 105, y0), radius=10, fill="#2563eb" if i == 0 else "#94a3b8")
        draw.text((bx + 12, y0 - bh - 28), f"{val:.4f}", font=value_font, fill="#0f172a")
        draw_wrapped(draw, (bx - 25, y0 + 22), labels[i], pil_font(21), "#0f172a", 160, line_gap=4, align="center")
    draw.text((1055, 845), "继续训练没有超过原始 best checkpoint", font=value_font, fill="#475569")
    draw.text((90, 930), "图源说明：作者根据 BUSI seed 稳定性实验与 CVC 继续训练实验结果整理绘制。", font=pil_font(20), fill="#475569")
    img.save(out, quality=95)
    return out


def _mask_path(dataset: str, sample_id: str) -> Path:
    if dataset == "busi":
        return ROOT / "data" / "processed" / "busi" / "masks" / "0" / f"{sample_id}_mask.png"
    return ROOT / "data" / "processed" / "cvc" / "masks" / "0" / f"{sample_id}.png"


def _image_path(dataset: str, sample_id: str) -> Path:
    return ROOT / "data" / "processed" / dataset / "images" / f"{sample_id}.png"


def _pred_path(run: str, sample_id: str) -> Path:
    return ROOT / "experiments" / "results" / run / "predictions" / f"{sample_id}.png"


def _binary(path: Path) -> Image.Image:
    return Image.open(path).convert("L").point(lambda x: 255 if x >= 128 else 0)


def dice_for_sample(gt: Image.Image, pred: Image.Image) -> float:
    g = gt.resize((256, 256), Image.Resampling.NEAREST)
    p = pred.resize((256, 256), Image.Resampling.NEAREST)
    gp = g.tobytes()
    pp = p.tobytes()
    inter = sum(1 for a, b in zip(gp, pp) if a > 0 and b > 0)
    total = sum(1 for a in gp if a > 0) + sum(1 for b in pp if b > 0)
    return 1.0 if total == 0 else 2 * inter / total


def error_overlay(image: Image.Image, gt: Image.Image, pred: Image.Image) -> Image.Image:
    base = image.convert("RGB").resize((256, 256), Image.Resampling.BILINEAR)
    gt = gt.resize((256, 256), Image.Resampling.NEAREST)
    pred = pred.resize((256, 256), Image.Resampling.NEAREST)
    out = base.copy()
    pix = out.load()
    gp = gt.load()
    pp = pred.load()
    for y in range(256):
        for x in range(256):
            g = gp[x, y] > 0
            p = pp[x, y] > 0
            if p and not g:
                r, g0, b = pix[x, y]
                pix[x, y] = (min(255, int(r * 0.55 + 255 * 0.45)), int(g0 * 0.55), int(b * 0.55))
            elif g and not p:
                r, g0, b = pix[x, y]
                pix[x, y] = (int(r * 0.55), int(g0 * 0.55), min(255, int(b * 0.55 + 255 * 0.45)))
            elif g and p:
                r, g0, b = pix[x, y]
                pix[x, y] = (int(r * 0.65), min(255, int(g0 * 0.65 + 255 * 0.35)), int(b * 0.65))
    return out


def create_error_case_figure() -> Path:
    out = FIG_DIR / "error_case_analysis.png"
    candidates = [
        ("BUSI", "busi", "busi_ukan_seed2981", ROOT / "data" / "splits" / "busi_seed2981_val.txt"),
        ("CVC", "cvc", "cvc_ukan_seed2981", ROOT / "data" / "splits" / "cvc_seed2981_val.txt"),
    ]
    selected: list[tuple[str, str, str, float]] = []
    for label, dataset, run, split_path in candidates:
        scored: list[tuple[str, float]] = []
        for sample_id in split_path.read_text(encoding="utf-8").splitlines():
            ip, mp, pp = _image_path(dataset, sample_id), _mask_path(dataset, sample_id), _pred_path(run, sample_id)
            if ip.exists() and mp.exists() and pp.exists():
                gt, pred = _binary(mp), _binary(pp)
                d = dice_for_sample(gt, pred)
                if 0.05 < d < 0.92:
                    scored.append((sample_id, d))
        scored.sort(key=lambda x: x[1])
        for sid, d in scored[:2]:
            selected.append((label, dataset, sid, d))
    selected = selected[:4]

    W, H = 1850, 1260
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "U-KAN 预测误差案例：边界偏差、漏分割与过分割", W, 35)
    headers = ["原图", "真值", "U-KAN 预测", "误差叠加"]
    font = pil_font(22, True)
    for j, h in enumerate(headers):
        draw.text((280 + j * 360, 120), h, font=font, fill="#0f172a")
    row_y = 170
    for i, (label, dataset, sid, d) in enumerate(selected):
        y = row_y + i * 260
        draw.text((55, y + 105), f"{label}\n{sid}\nDice={d:.3f}", font=pil_font(20), fill="#334155")
        image = Image.open(_image_path(dataset, sid)).convert("RGB").resize((220, 220), Image.Resampling.BILINEAR)
        gt = _binary(_mask_path(dataset, sid)).resize((220, 220), Image.Resampling.NEAREST).convert("RGB")
        pred = _binary(_pred_path(f"{dataset}_ukan_seed2981", sid)).resize((220, 220), Image.Resampling.NEAREST).convert("RGB")
        err = error_overlay(Image.open(_image_path(dataset, sid)), _binary(_mask_path(dataset, sid)), _binary(_pred_path(f"{dataset}_ukan_seed2981", sid))).resize((220, 220))
        for j, panel in enumerate([image, gt, pred, err]):
            x = 250 + j * 360
            draw.rounded_rectangle((x - 8, y - 8, x + 228, y + 228), radius=14, fill="#ffffff", outline="#cbd5e1", width=2)
            img.paste(panel, (x, y))
    draw.text((80, 1210), "图源说明：作者根据验证集原图、真值 mask 与 U-KAN 预测结果生成。误差叠加：绿色为命中区域，红色为过分割，蓝色为漏分割。", font=pil_font(19), fill="#475569")
    img.save(out, quality=95)
    return out


def create_workflow_v2() -> Path:
    out = FIG_DIR / "experiment_workflow.png"
    W, H = 2500, 1100
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "实验流程与课程论文产出链路", W, 42)
    body = pil_font(24)
    small = pil_font(20)
    items = [
        ("公开数据集\nBUSI / CVC", "#e0f2fe", "#0284c7", 250),
        ("数据整理\nmask 合并、二值化\n固定 train/val split", "#eef2ff", "#4f46e5", 310),
        ("四组模型\nU-Net / no-KAN\nU-KAN / Attn-U-KAN", "#ecfdf5", "#059669", 310),
        ("统一训练\n100 epoch\nBCE+Dice Loss", "#fef3c7", "#d97706", 250),
        ("统一评估\nIoU/Dice/Precision\nRecall/Specificity/耗时", "#fff7ed", "#ea580c", 310),
        ("补充实验\nBUSI seed6142\nCVC 继续训练", "#fce7f3", "#db2777", 250),
        ("论文与 GitHub\n图表、README\n复现命令、结论", "#f8fafc", "#64748b", 250),
    ]
    x = 90
    y = 355
    prev_end = None
    for text, fill, outline, w in items:
        draw_box(draw, (x, y, x + w, y + 160), text, fill, outline, body)
        if prev_end:
            arrow(draw, prev_end, (x, y + 80), "#64748b", 5)
        prev_end = (x + w, y + 80)
        x += w + 50
    draw_box(
        draw,
        (160, 720, 2340, 910),
        "关键控制：所有主实验使用同一输入分辨率 256、batch size 8、seed 2981、Adam 与 CosineAnnealingLR；最终量化结果统一来自 evaluate.py 的全验证集像素级混淆矩阵口径。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw.text((90, 1000), "图源说明：作者根据本文实验脚本、数据处理脚本和论文写作流程自绘。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def create_figures_v2() -> dict[str, Path]:
    figs = generate_figures()
    figs["workflow"] = create_workflow_v2()
    figs.update(copy_reference_figures())
    figs["dataset_combo"] = create_dataset_combo()
    figs["results_chart"] = create_results_chart()
    figs["diagnosis_flow"] = create_diagnosis_flow()
    figs["supplement_chart"] = create_supplement_chart()
    figs["error_cases"] = create_error_case_figure()
    return figs


def literature_rows() -> list[list[str]]:
    return [
        ["U-KAN", "Related Work、Method、Experiments、Ablation Studies", "SOTA 对比、效率表、定性分割图、KAN 层数消融、KAN vs MLP 消融", "本文采用 U-KAN/no-KAN/效率/可视化的基本实验骨架"],
        ["ResU-KAN", "Related work、Method、Experimental results and analysis", "Dataset、Loss、Metrics、Parameter、SOTA、Visualization、Boundary/error analysis、Ablation", "本文新增训练异常诊断、误差案例和补充实验讨论"],
        ["KC-UNet", "架构说明、预处理图、消融与多数据集对比", "KAN 层数、MLP 替代、attention 机制、四数据集结果、效率表、Dice 对比图", "本文把 no-KAN 和 Attention-U-KAN 写成明确消融，而非简单改进"],
        ["TransUKAN / VMKLA-UNet", "Related Work、Methodology、Implementation Details、SOTA、Ablation", "结构图、模块图、复杂度/推理时间、跨数据集可视化", "本文保留参数量与推理时间，讨论精度和效率权衡"],
        ["KAN 医学综述", "按任务和模型类型归纳 KAN 医学应用", "模型、任务、指标和代表性收益汇总表", "本文用它说明医学分割方向和 U-KAN 选题来源"],
    ]


def official_compare_rows() -> list[list[str]]:
    rows = load_results()
    return [
        ["BUSI U-KAN", "0.6526", "0.7875", fmt4(find_row(rows, "busi_ukan_seed2981")["iou"]), fmt4(find_row(rows, "busi_ukan_seed2981")["dice"]), "本地复现结果高于官方 README 单 run 参考值"],
        ["BUSI no-KAN", "0.6349", "0.7707", fmt4(find_row(rows, "busi_no_kan_seed2981")["iou"]), fmt4(find_row(rows, "busi_no_kan_seed2981")["dice"]), "趋势一致：U-KAN 高于 no-KAN"],
        ["CVC U-KAN", "0.8561", "0.9219", fmt4(find_row(rows, "cvc_ukan_seed2981")["iou"]), fmt4(find_row(rows, "cvc_ukan_seed2981")["dice"]), "低于官方参考，需结合 epoch、split、checkpoint 与多 seed 差异解释"],
    ]


def efficiency_rows() -> list[list[str]]:
    rows = load_results()
    names = [
        "busi_unet_seed2981",
        "busi_no_kan_seed2981",
        "busi_ukan_seed2981",
        "busi_attention_ukan_seed2981",
        "cvc_unet_seed2981",
        "cvc_no_kan_seed2981",
        "cvc_ukan_seed2981",
        "cvc_attention_ukan_seed2981",
    ]
    out = []
    for name in names:
        r = find_row(rows, name)
        out.append([r["dataset"].upper(), r["variant"], fmt_params(r["params"]), fmt2(r["infer_ms_per_image"]), fmt4(r["iou"]), fmt4(r["dice"])])
    return out


def ablation_summary_rows() -> list[list[str]]:
    rows = load_results()
    b_ukan = find_row(rows, "busi_ukan_seed2981")
    b_nokan = find_row(rows, "busi_no_kan_seed2981")
    b_att = find_row(rows, "busi_attention_ukan_seed2981")
    c_ukan = find_row(rows, "cvc_ukan_seed2981")
    c_nokan = find_row(rows, "cvc_no_kan_seed2981")
    c_att = find_row(rows, "cvc_attention_ukan_seed2981")
    return [
        ["KAN vs MLP(no-KAN)", "BUSI", f"+{float(b_ukan['iou'])-float(b_nokan['iou']):.4f}", f"+{float(b_ukan['dice'])-float(b_nokan['dice']):.4f}", "KANLinear 相比普通 Linear 提升明显"],
        ["KAN vs MLP(no-KAN)", "CVC", f"+{float(c_ukan['iou'])-float(c_nokan['iou']):.4f}", f"+{float(c_ukan['dice'])-float(c_nokan['dice']):.4f}", "提升较小但方向一致"],
        ["Attention 插入", "BUSI", f"{float(b_att['iou'])-float(b_ukan['iou']):.4f}", f"{float(b_att['dice'])-float(b_ukan['dice']):.4f}", "简单通道-空间注意力未超过 U-KAN"],
        ["Attention 插入", "CVC", f"{float(c_att['iou'])-float(c_ukan['iou']):.4f}", f"{float(c_att['dice'])-float(c_ukan['dice']):.4f}", "Precision 较高但 Recall 降低，预测更保守"],
    ]


def md_image(path: Path, caption: str, source: str) -> str:
    rel = path.relative_to(PAPER_DIR).as_posix() if path.is_relative_to(PAPER_DIR) else path.as_posix()
    return f"![{caption}]({rel})\n\n{caption}。来源：{source}\n"


def build_markdown_v2(figs: dict[str, Path], tables: dict[str, list[list[str]]], stats: dict[str, str]) -> None:
    headers = ["模型", "IoU", "Dice", "Precision", "Recall", "Specificity", "参数量", "推理 ms/图"]
    content = f"""# {TITLE}

**课程方向**：方向（一）复现任意人工智能算法  
**项目 GitHub**：[{GITHUB_URL}]({GITHUB_URL})  
**作者**：蔡雪峰  
**日期**：2026 年 6 月

## 来源说明与过程声明

本文在课程材料、U-KAN 官方论文与源码、KAN 医学图像分析综述、ResU-KAN、KC-UNet、TransUKAN、VMKLA-UNet、LightM-UNet 等论文基础上确定写作结构和实验组织方式。论文并非简单复述 U-KAN，而是在本地 PyTorch 工程中完成数据整理、模型复现、对比实验、异常曲线诊断和补充验证。所有实验表格均来自本项目 `evaluate.py` 和训练日志；所有训练曲线、预测对比、误差案例和截图均来自本地运行结果。

图源分为四类：第一，作者原始 PPT 手绘图，文件为 `{HANDDRAWN_PDF.name}`，本文将其作为 KAN 与 MLP 对比思路来源；第二，U-KAN 官方框架图和 KAN 医学综述表格，用于说明算法来源和选题背景，并在图注中标明出处；第三，本项目实验脚本生成的数据样本、训练曲线、预测结果和误差分析；第四，作者根据工程流程和实验思考过程整理绘制的实验流程图、诊断流程图和补充实验概览图。本文写作和排版使用 OpenAI Codex 辅助，但实验设计、运行过程、指标和结论均基于本地项目记录。

## 摘要

医学图像分割是计算机辅助诊断中的关键任务。U-Net 系列模型通过编码器、解码器和跳跃连接获得了稳定的分割性能，但普通卷积和 MLP 对复杂非线性边界的表达仍存在限制。U-KAN 将 Kolmogorov-Arnold Network 的可学习边函数引入 U 型结构，在 tokenized KAN block 中用样条函数增强非线性建模能力。本文选择 U-KAN 作为人工智能课程论文的复现对象，在 BUSI 乳腺超声和 CVC-ClinicDB 结肠镜息肉数据集上复现 U-KAN，并构建 U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN 四组对照。

与早期草稿不同，本文重点体现实验思考过程：在 BUSI 训练曲线中，no-KAN 的 loss 曲线较高但验证 IoU 局部跳升，引发“该结果是否偶然”的疑问，因此补充 seed 6142 稳定性实验；在 CVC 训练中，多个模型 best epoch 靠近训练末端，引发“是否未收敛”的疑问，因此从 U-KAN best checkpoint 继续训练 50 epoch。结果显示，U-KAN 在 BUSI 上取得 IoU={stats['busi_ukan_iou']}、Dice={stats['busi_ukan_dice']}，在 CVC 上取得 IoU={stats['cvc_ukan_iou']}、Dice={stats['cvc_ukan_dice']}，均为主实验最优；seed 6142 中 U-KAN 仍高于 no-KAN，但 BUSI 绝对指标对 seed 敏感；CVC 继续训练未超过原始 best checkpoint。实验说明 KAN 模块对医学图像分割有实际贡献，但简单注意力插入并未稳定提升，且 U-KAN 存在推理速度代价。

**关键词**：U-KAN；Kolmogorov-Arnold Network；医学图像分割；算法复现；消融实验；实验诊断

## 1 引言

医学图像分割要求模型在噪声强、目标边界模糊、标注数据有限的条件下定位病灶区域。BUSI 乳腺超声图像中的病灶边界常受声影、组织纹理和弱对比影响；CVC-ClinicDB 中息肉形态相对规则，但内镜反光、褶皱和边缘模糊仍会造成误分割。U-Net 是医学图像分割的经典框架，其编码器-解码器结构和跳跃连接适合恢复空间细节，但模型的中高层非线性表达通常依赖卷积或普通 MLP。

KAN 的核心思想是把可学习的一维函数放在网络边上，而不是只在节点使用固定激活函数。U-KAN 将这一思想嵌入 U 型医学图像分割网络，在高层 token 特征上使用 KAN layer 和 depthwise convolution 构成 tokenized KAN block。与直接引入 Transformer 相比，U-KAN 的动机更强调非线性函数表达和一定可解释性。课程论文选择 U-KAN 的原因有三点：其一，它是明确的人工智能算法复现对象；其二，它天然适合医学图像分割实验，包括数据整理、对比实验和可视化；其三，KAN vs MLP、U-KAN vs U-Net、attention 改进是否有效等问题可以形成完整的消融分析。

本文贡献如下。第一，参考 U-KAN 官方代码重写并组织课程工程，完成 BUSI/CVC 数据处理、训练、评估和结果汇总。第二，构建 U-Net、no-KAN、U-KAN 和 Attention-U-KAN 四组模型，对 KAN 模块和注意力模块进行可解释消融。第三，针对训练曲线中出现的异常现象提出假设并补做实验，体现从结果观察到实验验证的研究过程。第四，整理 GitHub、README、运行截图、自绘图、论文图源和完整 Word 成稿，覆盖课程论文要求和加分项。

## 2 相关工作与论文实验组织参考

为了避免论文只停留在“跑代码和贴结果”，本文先阅读并整理了 U-KAN 及多篇 KAN 医学图像分割论文的实验组织方式。U-KAN 原论文包含 SOTA 对比、效率对比、定性分割图、可解释性图、KAN 层数消融、KAN vs MLP 消融和模型规模消融。ResU-KAN 的实验章节更细，包括数据集、损失函数、评价指标、参数设置、SOTA 对比、可视化分析、边界误差分析和消融实验。KC-UNet 则专门比较 KAN 层数、MLP 替代、不同 attention 机制和模型效率。这些论文共同说明，医学图像分割论文不能只报告一个主表，而应同时给出结构图、数据图、训练与预测可视化、消融、效率和局限讨论。

{markdown_table(['参考论文', '文章组织', '实验/图表特点', '对本文的借鉴'], literature_rows())}

表 1 来源说明：作者根据本地课程材料中的 U-KAN、ResU-KAN、KC-UNet、TransUKAN、VMKLA-UNet 与 KAN 医学综述论文整理。

{md_image(figs['kan_medical_review_table'], '图 1 KAN 医学图像分析综述中的代表模型与任务汇总', 'A Review of KANs for Medical Image Analysis，本地课程材料。本文用于说明 KAN 医学应用背景。')}

## 3 方法

### 3.1 KAN 与 MLP 的区别

传统 MLP 通常学习线性权重，并在节点使用固定激活函数，如 ReLU、GELU 或 SiLU。KAN 则将每条边的映射定义为可学习的一维函数，可用 B-spline 基函数展开：φ(x)=ΣciBi(x)。这种形式使网络不只依赖固定节点激活，而是在边上学习更灵活的非线性映射。作者的 PPT 手绘图最初正是围绕这一差异展开：MLP 是 fixed node activations + scalar edge weights，而 KAN 是 learnable 1D edge functions。

{md_image(figs['kan_vs_mlp'], '图 2 MLP 与 KAN 的建模差异', f'作者根据 PPT 自绘图 {HANDDRAWN_PDF.name} 重新整理为论文位图；原始 PDF 已保存在 paper 目录。')}

### 3.2 U-KAN 复现模型

U-KAN 保留 U-Net 的编码器-解码器和 skip connection，但在中高层特征处引入 tokenized KAN block。低层卷积负责提取局部纹理，高层 PatchEmbed 将特征转为 token，KAN layer 对 token 进行非线性映射，再通过 depthwise convolution 加入局部空间归纳偏置。本文复现代码中，`U-KAN(no-KAN)` 保持同一 U 型结构，仅将 KANLinear 替换为普通 Linear，因此它不是 U-Net，而是用于隔离 KAN 模块贡献的消融模型。

{md_image(figs['official_ukan_framework'], '图 3 U-KAN 原论文/官方项目中的总体结构图', 'U-KAN 官方开源项目 assets/framework-1.jpg，本文用于说明复现对象的原始结构。')}

### 3.3 Attention-U-KAN 改进尝试

本文尝试在 skip fusion 后加入轻量通道-空间注意力模块。其动机是：skip connection 会把低层纹理和边界细节传给解码器，但也可能带入背景噪声；通道注意力可以重标定语义通道，空间注意力可以突出疑似病灶区域。然而，这只是一个可检验假设，不能预设一定提升。实验结果最终显示 Attention-U-KAN 没有超过原始 U-KAN，这一负结果本身也是本文消融分析的重要部分。

{md_image(figs['attention_module'], '图 4 Attention-U-KAN 的 skip fusion 增强模块', '作者根据本文 attention.py 与 attention_ukan.py 代码自行绘制。')}

## 4 实验设计

### 4.1 数据集与预处理

本文使用 BUSI 和 CVC-ClinicDB 两个公开数据集。BUSI 原始数据中部分图像有多个 mask，本文将同一图像的多个 mask 取并集合并，得到单通道二值病灶标注；CVC-ClinicDB 图像与 mask 一一对应。所有图像统一整理到 `data/processed`，并通过固定 split 文件保存训练/验证划分，避免不同模型使用不同样本划分。

{markdown_table(['数据集', '图像数', '宽度范围', '高度范围', '平均 mask 占比', '处理说明'], tables['dataset'])}

表 2 来源说明：由本文 `dataset_report.py` 对整理后的本地数据统计生成。

{md_image(figs['dataset_combo'], '图 5 BUSI 与 CVC-ClinicDB 数据样本可视化', '作者根据本文数据处理脚本生成。')}

### 4.2 实验设置与评价指标

所有主实验使用 256×256 输入分辨率、batch size 8、100 epoch、Adam 优化器和 CosineAnnealingLR。损失函数为 `0.5×BCEWithLogits + DiceLoss`。最终结果不使用训练日志中的 batch 平均指标，而统一使用 `evaluate.py` 在全验证集像素级混淆矩阵上计算 IoU、Dice、Precision、Recall 和 Specificity。这样可以避免训练曲线局部波动被误读为最终模型能力。

{markdown_table(['项目', '设置', '说明'], tables['platform'])}

表 3 来源说明：由 `doctor_env.py`、YAML 配置与训练脚本记录整理。

为了让复现过程可检查，本文将数据处理、训练、评估、绘图、补充实验和论文生成都落到脚本层面，形成从原始数据到课程论文材料的闭环。

{md_image(figs['workflow'], '图 6 项目实验流程与代码组织', '作者根据本文工程目录、训练脚本和论文生成流程自行绘制。')}

## 5 主实验结果

### 5.1 BUSI 结果

{markdown_table(headers, tables['busi_results'])}

表 4 来源说明：由 `evaluate.py` 在 BUSI 验证集上统一评估生成。

BUSI 上，U-KAN 取得 IoU={stats['busi_ukan_iou']}、Dice={stats['busi_ukan_dice']}，优于 U-Net、no-KAN 和 Attention-U-KAN。与 no-KAN 相比，U-KAN 的 IoU 提升 {stats['busi_gain_nokan_iou']}，说明在同一 U 型结构中引入 KANLinear 对病灶重叠质量有实际贡献。no-KAN 的 Precision 较高但 Recall 较低，说明其预测更保守；U-KAN 的 Recall 明显提升，更有利于减少漏分割。

{md_image(figs['busi_curves'], '图 7 BUSI 训练曲线', '作者根据训练日志生成。')}

{md_image(figs['busi_preds'], '图 8 BUSI 预测结果对比', '作者根据验证集预测结果生成。')}

### 5.2 CVC 结果

{markdown_table(headers, tables['cvc_results'])}

表 5 来源说明：由 `evaluate.py` 在 CVC 验证集上统一评估生成。

CVC 上，U-KAN 同样取得主实验最优 IoU 和 Dice。相比 BUSI，U-KAN 对 U-Net 的提升幅度较小，说明 CVC 数据中传统 U-Net 已经是较强基线。Attention-U-KAN 的 Precision 较高但 Recall 下降，表明该模型更容易做保守预测，漏掉部分真实息肉区域。

{md_image(figs['cvc_curves'], '图 9 CVC 训练曲线', '作者根据训练日志生成。')}

{md_image(figs['cvc_preds'], '图 10 CVC 预测结果对比', '作者根据验证集预测结果生成。')}

{md_image(figs['results_chart'], '图 11 主实验 IoU/Dice 可视化对比', '作者根据 evaluate.py 指标整理绘制。')}

## 6 消融实验、异常诊断与补充验证

### 6.1 KAN 与注意力模块消融

{markdown_table(['消融问题', '数据集', 'IoU 变化', 'Dice 变化', '结论'], ablation_summary_rows())}

表 6 来源说明：由 U-KAN、no-KAN 与 Attention-U-KAN 的主实验结果整理。

U-KAN 相对 no-KAN 在 BUSI 和 CVC 上均有提升，因此 KAN 模块的贡献比单纯网络结构更可信。Attention-U-KAN 没有超过原始 U-KAN，说明“加注意力”不是必然加分项；更合理的结论是，本文尝试的简单通道-空间注意力插入方式未与 U-KAN 形成稳定互补。

### 6.2 从异常曲线到补充实验

训练过程中，BUSI no-KAN 的橙色 loss 曲线明显高于其他模型，下降也更慢，但验证 IoU/Dice 在局部 epoch 出现跳升。这一现象不应被直接解释为 no-KAN 学得更好，因为 BCE+Dice loss 是连续概率空间中的训练目标，而 IoU/Dice 是阈值化后的离散分割指标；当目标区域较小、验证集规模有限时，少数样本的阈值变化可能造成指标跳变。因此本文提出第一个疑问：no-KAN 的局部高点是否可能是随机划分或阈值化造成的偶然结果？

CVC 训练曲线中，部分模型的 best epoch 靠近 100 epoch 末端，同时本地 CVC U-KAN 低于官方 README 参考值，这引出第二个疑问：模型是否只是训练不够？围绕这两个疑问，本文设计了两个补充实验，而不是直接接受单次曲线。

{md_image(figs['diagnosis_flow'], '图 12 从异常曲线到补充实验的思考过程', '作者根据本项目实验讨论和补充实验流程整理绘制。')}

### 6.3 BUSI seed 稳定性实验

{markdown_table(['模型', 'seed2981 IoU', 'seed2981 Dice', 'seed6142 IoU', 'seed6142 Dice', '两 seed 平均 IoU', '两 seed 平均 Dice'], tables['seed_stability'])}

表 7 来源说明：由 BUSI seed 2981 与 seed 6142 的 `evaluate.py` 结果整理。

seed 6142 下，U-KAN 仍高于 no-KAN；两 seed 平均后，U-KAN IoU 为 0.6515，高于 no-KAN 的 0.6214。这说明 KAN 的收益不是 seed 2981 的单次偶然峰值。但 seed 6142 的绝对指标整体下降，也说明 BUSI 对数据划分和初始化敏感。因此论文结论必须谨慎：本文证明了本地复现条件下 KAN 有稳定趋势，但不应夸大单次 seed 的数值。

### 6.4 CVC 继续训练实验

{markdown_table(['实验', '设置', 'epoch', 'best epoch', 'IoU', 'Dice', '后 10 epoch IoU 变化'], tables['cvc_continue'])}

表 8 来源说明：由 CVC U-KAN 原始训练与继续训练结果整理。

从 CVC U-KAN best checkpoint 继续训练 50 epoch 后，IoU 为 0.7847，未超过原始 100 epoch 的 0.7874。因此，目前证据不支持“CVC 结果偏低只是因为没有继续训练”的解释。更可能的原因包括官方训练轮数更长、数据划分和预处理不同、官方 checkpoint 或多 seed 平均差异。

{md_image(figs['supplement_chart'], '图 13 补充实验结果概览', '作者根据 BUSI seed 稳定性与 CVC 继续训练实验整理绘制。')}

### 6.5 误差案例分析

{md_image(figs['error_cases'], '图 14 U-KAN 预测误差案例分析', '作者根据验证集原图、真值 mask 与 U-KAN 预测结果生成。')}

误差案例显示，U-KAN 的主要问题不是完全无法定位目标，而是在弱边界、小目标和边缘不规则区域容易出现漏分割或过分割。BUSI 中边界模糊和声影会让模型低估病灶范围；CVC 中反光和褶皱可能造成边界偏移。这也解释了为什么简单 attention 可能提高 Precision 却降低 Recall：注意力模块更保守地抑制不确定区域，反而漏掉部分真实目标。

## 7 与官方结果、效率和课程加分项的关系

{markdown_table(['项目', '官方 IoU', '官方 F1/Dice', '本文 IoU', '本文 Dice', '说明'], official_compare_rows())}

表 9 来源说明：官方参考值来自 U-KAN README；本文结果来自本地 `evaluate.py`。

{markdown_table(['数据集', '模型', '参数量', '推理 ms/图', 'IoU', 'Dice'], efficiency_rows())}

表 10 来源说明：参数量由模型统计脚本输出，推理时间由 `evaluate.py` 记录。

U-Net 的推理速度显著快于 U-KAN 系列，而 U-KAN 在 BUSI 和 CVC 上取得更好分割指标。这说明 KAN 带来的不是免费提升，而是精度和效率之间的权衡。课程论文中必须把这一点写清楚，否则会显得只选择性报告好结果。

{markdown_table(['要求/加分项', '本文落实方式', '位置'], [
['GitHub 源码链接', f'完整项目远程仓库 {GITHUB_URL}', '封面、来源说明、README'],
['详细 README', '环境、数据准备、训练、评估、绘图和论文生成命令', 'README.md'],
['项目结构', 'configs/scripts/src/tests/experiments/paper 分层组织', 'GitHub 仓库'],
['数据整理', 'BUSI 多 mask 合并、CVC 标准化、固定 split', '数据集章节与脚本'],
['自绘/作者图', 'PPT 手绘图、实验流程图、诊断流程图、补充实验图', '图 2、图 6、图 12、图 13'],
['论文图源规范', '每张图均注明作者自绘、论文引用或项目生成', '所有图注'],
['过程记录', '训练截图、曲线异常、补充实验命令与结果', '第 6 章与过程截图'],
['独立思考', '从 no-KAN 跳变和 CVC 未收敛疑问出发补做实验', '第 6.2–6.4 节'],
] )}

表 11 来源说明：根据课程论文要求和本文项目材料整理。

## 8 结论

本文按照正式医学图像分割论文的写法重新组织了 U-KAN 课程复现工作。实验表明，U-KAN 在 BUSI 和 CVC 主实验中均取得最优 IoU/Dice；no-KAN 消融说明 KANLinear 相比普通 Linear 具有实际贡献；Attention-U-KAN 未超过原始 U-KAN，说明简单注意力插入不是稳定改进方向。更重要的是，本文没有盲目接受单次结果，而是针对 BUSI no-KAN 的 loss 与指标矛盾、局部跳升，以及 CVC best epoch 靠后的现象提出疑问并补做实验。

补充实验显示，seed 6142 下 U-KAN 仍优于 no-KAN，但 BUSI 绝对指标受随机划分影响明显；CVC 继续训练未超过原始 best checkpoint，说明本地结果与官方参考差距不能简单归因于训练不够。最终，本文形成了包含源码、README、数据处理、训练日志、主实验、消融实验、补充实验、图表来源和 Word 成稿的完整课程论文项目，既满足算法复现要求，也体现了额外实践投入和独立分析过程。

## 参考文献

[1] Li C., Liu X., Li W., et al. U-KAN Makes Strong Backbone for Medical Image Segmentation and Generation. arXiv, 2024.  
[2] Liu Z., Wang Y., Vaidya S., et al. KAN: Kolmogorov-Arnold Networks. arXiv:2404.19756, 2024.  
[3] Ronneberger O., Fischer P., Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation. MICCAI, 2015.  
[4] Woo S., Park J., Lee J.-Y., Kweon I. S. CBAM: Convolutional Block Attention Module. ECCV, 2018.  
[5] Wang et al. ResU-KAN: U-KAN based residual model for medical image segmentation. 本地课程材料。  
[6] Xu et al. KC-UNet: U-Net with Tokenized KAN modules and attention for medical image segmentation. 本地课程材料。  
[7] TransUKAN: Computing-Efficient Hybrid KAN-Transformer for Enhanced Medical Image Segmentation. 本地课程材料。  
[8] VMKLA-UNet: Vision Mamba with KAN Linear Attention U-Net. 本地课程材料。  
[9] A Review of KANs for Medical Image Analysis. 本地课程材料。  
[10] Al-Dhabyani W., Gomaa M., Khaled H., Fahmy A. Dataset of Breast Ultrasound Images. Data in Brief, 2020.  
[11] Bernal J., Sánchez F. J., Fernández-Esparrach G., et al. WM-DOVA maps for accurate polyp highlighting in colonoscopy. Computerized Medical Imaging and Graphics, 2015.
"""
    MD_OUT.write_text(content, encoding="utf-8")


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run(TITLE)
    set_run_font(r, 22, True, east_asia="黑体")
    for item in [
        "课程方向：方向（一）复现任意人工智能算法",
        "课程：人工智能",
        "作者：蔡雪峰",
        "日期：2026 年 6 月",
        f"项目 GitHub：{GITHUB_URL}",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(item)
        set_run_font(run, 12)
    doc.add_page_break()


def init_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.35)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.45)
    sec.right_margin = Cm(2.25)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    return doc


def build_docx_v2(figs: dict[str, Path], tables: dict[str, list[list[str]]], stats: dict[str, str]) -> None:
    doc = init_doc()
    title_page(doc)
    headers = ["模型", "IoU", "Dice", "Precision", "Recall", "Specificity", "参数量", "推理 ms/图"]

    add_heading(doc, "来源说明与过程声明", 1)
    add_para(doc, f"本文在课程材料、U-KAN 官方论文与源码、KAN 医学图像分析综述、ResU-KAN、KC-UNet、TransUKAN、VMKLA-UNet、LightM-UNet 等论文基础上确定写作结构和实验组织方式。论文并非简单复述 U-KAN，而是在本地 PyTorch 工程中完成数据整理、模型复现、对比实验、异常曲线诊断和补充验证。所有实验表格均来自本项目 evaluate.py 和训练日志；所有训练曲线、预测对比、误差案例和截图均来自本地运行结果。")
    add_para(doc, f"图源分为四类：作者原始 PPT 手绘图 {HANDDRAWN_PDF.name}；U-KAN 官方框架图和 KAN 医学综述表格；本项目实验脚本生成的数据样本、训练曲线、预测结果和误差分析；作者根据工程流程和实验思考过程整理绘制的实验流程图、诊断流程图和补充实验概览图。本文写作和排版使用 OpenAI Codex 辅助，但实验设计、运行过程、指标和结论均基于本地项目记录。")

    add_heading(doc, "摘要", 1)
    add_para(doc, f"医学图像分割是计算机辅助诊断中的关键任务。U-Net 系列模型通过编码器、解码器和跳跃连接获得了稳定的分割性能，但普通卷积和 MLP 对复杂非线性边界的表达仍存在限制。U-KAN 将 Kolmogorov-Arnold Network 的可学习边函数引入 U 型结构，在 tokenized KAN block 中用样条函数增强非线性建模能力。本文选择 U-KAN 作为人工智能课程论文的复现对象，在 BUSI 乳腺超声和 CVC-ClinicDB 结肠镜息肉数据集上复现 U-KAN，并构建 U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN 四组对照。实验过程中，BUSI no-KAN 出现 loss 较高但验证 IoU 局部跳升的现象，CVC 多个模型 best epoch 靠近训练末端。针对这些疑问，本文补充 seed 6142 稳定性实验和 CVC 继续训练实验。结果显示，U-KAN 在 BUSI 上取得 IoU={stats['busi_ukan_iou']}、Dice={stats['busi_ukan_dice']}，在 CVC 上取得 IoU={stats['cvc_ukan_iou']}、Dice={stats['cvc_ukan_dice']}，均为主实验最优；Attention-U-KAN 未稳定超过原始 U-KAN。")
    add_para(doc, "关键词：U-KAN；Kolmogorov-Arnold Network；医学图像分割；算法复现；消融实验；实验诊断", first_line=False)

    add_heading(doc, "1 引言", 1)
    for text in [
        "医学图像分割要求模型在噪声强、目标边界模糊、标注数据有限的条件下定位病灶区域。BUSI 乳腺超声图像中的病灶边界常受声影、组织纹理和弱对比影响；CVC-ClinicDB 中息肉形态相对规则，但内镜反光、褶皱和边缘模糊仍会造成误分割。U-Net 是医学图像分割的经典框架，其编码器-解码器结构和跳跃连接适合恢复空间细节，但模型的中高层非线性表达通常依赖卷积或普通 MLP。",
        "KAN 的核心思想是把可学习的一维函数放在网络边上，而不是只在节点使用固定激活函数。U-KAN 将这一思想嵌入 U 型医学图像分割网络，在高层 token 特征上使用 KAN layer 和 depthwise convolution 构成 tokenized KAN block。课程论文选择 U-KAN 的原因有三点：它是明确的人工智能算法复现对象；它天然适合医学图像分割实验；KAN vs MLP、U-KAN vs U-Net、attention 改进是否有效等问题可以形成完整的消融分析。",
        "本文贡献包括：参考 U-KAN 官方代码重写课程工程；构建 U-Net、no-KAN、U-KAN 和 Attention-U-KAN 对照；针对训练曲线异常提出假设并补做实验；整理 GitHub、README、运行截图、自绘图、论文图源和完整 Word 成稿，覆盖课程论文要求和加分项。",
    ]:
        add_para(doc, text)

    add_heading(doc, "2 相关工作与论文实验组织参考", 1)
    add_para(doc, "为了避免论文只停留在跑代码和贴结果，本文先阅读并整理了 U-KAN 及多篇 KAN 医学图像分割论文的实验组织方式。U-KAN 原论文包含 SOTA 对比、效率对比、定性分割图、可解释性图、KAN 层数消融、KAN vs MLP 消融和模型规模消融。ResU-KAN 的实验章节更细，包括数据集、损失函数、评价指标、参数设置、SOTA 对比、可视化分析、边界误差分析和消融实验。KC-UNet 则专门比较 KAN 层数、MLP 替代、不同 attention 机制和模型效率。")
    add_table(doc, "表 1 相关论文实验组织方式对本文的启发", ["参考论文", "文章组织", "实验/图表特点", "对本文的借鉴"], literature_rows(), "作者根据本地课程材料中的相关论文整理。", 7.6)
    add_figure(doc, figs["kan_medical_review_table"], "图 1 KAN 医学图像分析综述中的代表模型与任务汇总", "A Review of KANs for Medical Image Analysis，本地课程材料。本文用于说明 KAN 医学应用背景。", 15.8)

    add_heading(doc, "3 方法", 1)
    add_heading(doc, "3.1 KAN 与 MLP 的区别", 2)
    add_para(doc, "传统 MLP 通常学习线性权重，并在节点使用固定激活函数，如 ReLU、GELU 或 SiLU。KAN 则将每条边的映射定义为可学习的一维函数，可用 B-spline 基函数展开：φ(x)=ΣciBi(x)。作者的 PPT 手绘图最初正是围绕这一差异展开：MLP 是 fixed node activations + scalar edge weights，而 KAN 是 learnable 1D edge functions。")
    add_figure(doc, figs["kan_vs_mlp"], "图 2 MLP 与 KAN 的建模差异", f"作者根据 PPT 自绘图 {HANDDRAWN_PDF.name} 重新整理为论文位图；原始 PDF 已保存在 paper 目录。", 15.8)
    add_heading(doc, "3.2 U-KAN 复现模型", 2)
    add_para(doc, "U-KAN 保留 U-Net 的编码器-解码器和 skip connection，但在中高层特征处引入 tokenized KAN block。低层卷积负责提取局部纹理，高层 PatchEmbed 将特征转为 token，KAN layer 对 token 进行非线性映射，再通过 depthwise convolution 加入局部空间归纳偏置。本文复现代码中，U-KAN(no-KAN) 保持同一 U 型结构，仅将 KANLinear 替换为普通 Linear，因此它不是 U-Net，而是用于隔离 KAN 模块贡献的消融模型。")
    add_figure(doc, figs["official_ukan_framework"], "图 3 U-KAN 原论文/官方项目中的总体结构图", "U-KAN 官方开源项目 assets/framework-1.jpg，本文用于说明复现对象的原始结构。", 16.0)
    add_heading(doc, "3.3 Attention-U-KAN 改进尝试", 2)
    add_para(doc, "本文尝试在 skip fusion 后加入轻量通道-空间注意力模块。其动机是：skip connection 会把低层纹理和边界细节传给解码器，但也可能带入背景噪声；通道注意力可以重标定语义通道，空间注意力可以突出疑似病灶区域。然而，这只是一个可检验假设，不能预设一定提升。")
    add_figure(doc, figs["attention_module"], "图 4 Attention-U-KAN 的 skip fusion 增强模块", "作者根据本文 attention.py 与 attention_ukan.py 代码自行绘制。", 15.8)

    add_heading(doc, "4 实验设计", 1)
    add_para(doc, "本文使用 BUSI 和 CVC-ClinicDB 两个公开数据集。BUSI 原始数据中部分图像有多个 mask，本文将同一图像的多个 mask 取并集合并，得到单通道二值病灶标注；CVC-ClinicDB 图像与 mask 一一对应。所有图像统一整理到 data/processed，并通过固定 split 文件保存训练/验证划分，避免不同模型使用不同样本划分。")
    add_table(doc, "表 2 数据集统计与预处理说明", ["数据集", "图像数", "宽度范围", "高度范围", "平均 mask 占比", "处理说明"], tables["dataset"], "由本文 dataset_report.py 对整理后的本地数据统计生成。", 8.0)
    add_figure(doc, figs["dataset_combo"], "图 5 BUSI 与 CVC-ClinicDB 数据样本可视化", "作者根据本文数据处理脚本生成。", 15.8)
    add_para(doc, "所有主实验使用 256×256 输入分辨率、batch size 8、100 epoch、Adam 优化器和 CosineAnnealingLR。损失函数为 0.5×BCEWithLogits + DiceLoss。最终结果统一使用 evaluate.py 在全验证集像素级混淆矩阵上计算指标。")
    add_table(doc, "表 3 实验平台与统一训练设置", ["项目", "设置", "说明"], tables["platform"], "由 doctor_env.py、YAML 配置与训练脚本记录整理。", 9.0)
    add_para(doc, "为了让复现过程可检查，本文将数据处理、训练、评估、绘图、补充实验和论文生成都落到脚本层面，形成从原始数据到课程论文材料的闭环。")
    add_figure(doc, figs["workflow"], "图 6 项目实验流程与代码组织", "作者根据本文工程目录、训练脚本和论文生成流程自行绘制。", 15.8)

    add_heading(doc, "5 主实验结果", 1)
    add_heading(doc, "5.1 BUSI 结果", 2)
    add_table(doc, "表 4 BUSI 主实验结果", headers, tables["busi_results"], "由 evaluate.py 在 BUSI 验证集上统一评估生成。", 7.6)
    add_para(doc, f"BUSI 上，U-KAN 取得 IoU={stats['busi_ukan_iou']}、Dice={stats['busi_ukan_dice']}，优于 U-Net、no-KAN 和 Attention-U-KAN。与 no-KAN 相比，U-KAN 的 IoU 提升 {stats['busi_gain_nokan_iou']}，说明在同一 U 型结构中引入 KANLinear 对病灶重叠质量有实际贡献。")
    add_figure(doc, figs["busi_curves"], "图 7 BUSI 训练曲线", "作者根据训练日志生成。", 15.8)
    add_figure(doc, figs["busi_preds"], "图 8 BUSI 预测结果对比", "作者根据验证集预测结果生成。", 15.8)
    add_heading(doc, "5.2 CVC 结果", 2)
    add_table(doc, "表 5 CVC 主实验结果", headers, tables["cvc_results"], "由 evaluate.py 在 CVC 验证集上统一评估生成。", 7.6)
    add_para(doc, f"CVC 上，U-KAN 同样取得主实验最优 IoU 和 Dice，分别为 {stats['cvc_ukan_iou']} 和 {stats['cvc_ukan_dice']}。相比 BUSI，U-KAN 对 U-Net 的提升幅度较小，说明 CVC 数据中传统 U-Net 已经是较强基线。")
    add_figure(doc, figs["cvc_curves"], "图 9 CVC 训练曲线", "作者根据训练日志生成。", 15.8)
    add_figure(doc, figs["cvc_preds"], "图 10 CVC 预测结果对比", "作者根据验证集预测结果生成。", 15.8)
    add_figure(doc, figs["results_chart"], "图 11 主实验 IoU/Dice 可视化对比", "作者根据 evaluate.py 指标整理绘制。", 15.8)

    add_heading(doc, "6 消融实验、异常诊断与补充验证", 1)
    add_table(doc, "表 6 KAN 与注意力模块消融总结", ["消融问题", "数据集", "IoU 变化", "Dice 变化", "结论"], ablation_summary_rows(), "由 U-KAN、no-KAN 与 Attention-U-KAN 的主实验结果整理。", 8.0)
    add_para(doc, "训练过程中，BUSI no-KAN 的橙色 loss 曲线明显高于其他模型，下降也更慢，但验证 IoU/Dice 在局部 epoch 出现跳升。这一现象不应被直接解释为 no-KAN 学得更好，因为 BCE+Dice loss 是连续概率空间中的训练目标，而 IoU/Dice 是阈值化后的离散分割指标；当目标区域较小、验证集规模有限时，少数样本的阈值变化可能造成指标跳变。")
    add_para(doc, "CVC 训练曲线中，部分模型的 best epoch 靠近 100 epoch 末端，同时本地 CVC U-KAN 低于官方 README 参考值，这引出第二个疑问：模型是否只是训练不够？围绕这两个疑问，本文设计了两个补充实验，而不是直接接受单次曲线。")
    add_figure(doc, figs["diagnosis_flow"], "图 12 从异常曲线到补充实验的思考过程", "作者根据本项目实验讨论和补充实验流程整理绘制。", 15.8)
    add_table(doc, "表 7 BUSI seed 稳定性补充实验", ["模型", "seed2981 IoU", "seed2981 Dice", "seed6142 IoU", "seed6142 Dice", "两 seed 平均 IoU", "两 seed 平均 Dice"], tables["seed_stability"], "由 BUSI seed 2981 与 seed 6142 的 evaluate.py 结果整理。", 7.6)
    add_para(doc, "seed 6142 下，U-KAN 仍高于 no-KAN；两 seed 平均后，U-KAN IoU 为 0.6515，高于 no-KAN 的 0.6214。这说明 KAN 的收益不是 seed 2981 的单次偶然峰值。但 seed 6142 的绝对指标整体下降，也说明 BUSI 对数据划分和初始化敏感。")
    add_table(doc, "表 8 CVC U-KAN 继续训练补充实验", ["实验", "设置", "epoch", "best epoch", "IoU", "Dice", "后 10 epoch IoU 变化"], tables["cvc_continue"], "由 CVC U-KAN 原始训练与继续训练结果整理。", 7.6)
    add_para(doc, "从 CVC U-KAN best checkpoint 继续训练 50 epoch 后，IoU 为 0.7847，未超过原始 100 epoch 的 0.7874。因此，目前证据不支持“CVC 结果偏低只是因为没有继续训练”的解释。")
    add_figure(doc, figs["supplement_chart"], "图 13 补充实验结果概览", "作者根据 BUSI seed 稳定性与 CVC 继续训练实验整理绘制。", 15.8)
    add_figure(doc, figs["error_cases"], "图 14 U-KAN 预测误差案例分析", "作者根据验证集原图、真值 mask 与 U-KAN 预测结果生成。", 15.8)
    add_para(doc, "误差案例显示，U-KAN 的主要问题不是完全无法定位目标，而是在弱边界、小目标和边缘不规则区域容易出现漏分割或过分割。BUSI 中边界模糊和声影会让模型低估病灶范围；CVC 中反光和褶皱可能造成边界偏移。")

    add_heading(doc, "7 与官方结果、效率和课程加分项的关系", 1)
    add_table(doc, "表 9 与 U-KAN 官方参考结果的对照", ["项目", "官方 IoU", "官方 F1/Dice", "本文 IoU", "本文 Dice", "说明"], official_compare_rows(), "官方参考值来自 U-KAN README；本文结果来自本地 evaluate.py。", 7.6)
    add_table(doc, "表 10 参数量、推理效率与指标对照", ["数据集", "模型", "参数量", "推理 ms/图", "IoU", "Dice"], efficiency_rows(), "参数量由模型统计脚本输出，推理时间由 evaluate.py 记录。", 7.6)
    add_para(doc, "U-Net 的推理速度显著快于 U-KAN 系列，而 U-KAN 在 BUSI 和 CVC 上取得更好分割指标。这说明 KAN 带来的不是免费提升，而是精度和效率之间的权衡。课程论文中必须把这一点写清楚，否则会显得只选择性报告好结果。")
    bonus_rows = [
        ["GitHub 源码链接", f"完整项目远程仓库 {GITHUB_URL}", "封面、来源说明、README"],
        ["详细 README", "环境、数据准备、训练、评估、绘图和论文生成命令", "README.md"],
        ["项目结构", "configs/scripts/src/tests/experiments/paper 分层组织", "GitHub 仓库"],
        ["数据整理", "BUSI 多 mask 合并、CVC 标准化、固定 split", "数据集章节与脚本"],
        ["自绘/作者图", "PPT 手绘图、实验流程图、诊断流程图、补充实验图", "图 2、图 6、图 12、图 13"],
        ["论文图源规范", "每张图均注明作者自绘、论文引用或项目生成", "所有图注"],
        ["过程记录", "训练截图、曲线异常、补充实验命令与结果", "第 6 章与过程截图"],
        ["独立思考", "从 no-KAN 跳变和 CVC 未收敛疑问出发补做实验", "第 6.2–6.4 节"],
    ]
    add_table(doc, "表 11 课程要求与加分项落实情况", ["要求/加分项", "本文落实方式", "位置"], bonus_rows, "根据课程论文要求和本文项目材料整理。", 7.8)

    add_heading(doc, "8 结论", 1)
    add_para(doc, "本文按照正式医学图像分割论文的写法重新组织了 U-KAN 课程复现工作。实验表明，U-KAN 在 BUSI 和 CVC 主实验中均取得最优 IoU/Dice；no-KAN 消融说明 KANLinear 相比普通 Linear 具有实际贡献；Attention-U-KAN 未超过原始 U-KAN，说明简单注意力插入不是稳定改进方向。更重要的是，本文没有盲目接受单次结果，而是针对 BUSI no-KAN 的 loss 与指标矛盾、局部跳升，以及 CVC best epoch 靠后的现象提出疑问并补做实验。")
    add_para(doc, "补充实验显示，seed 6142 下 U-KAN 仍优于 no-KAN，但 BUSI 绝对指标受随机划分影响明显；CVC 继续训练未超过原始 100 epoch best checkpoint，说明本地 U-KAN 结果与官方 CVC 参考差距不能简单归因于训练不够。最终，本文形成了包含源码、README、数据处理、训练日志、主实验、消融实验、补充实验、图表来源和 Word 成稿的完整课程论文项目，既满足算法复现要求，也体现了额外实践投入和独立分析过程。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "Li C., Liu X., Li W., et al. U-KAN Makes Strong Backbone for Medical Image Segmentation and Generation. arXiv, 2024.",
        "Liu Z., Wang Y., Vaidya S., et al. KAN: Kolmogorov-Arnold Networks. arXiv:2404.19756, 2024.",
        "Ronneberger O., Fischer P., Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation. MICCAI, 2015.",
        "Woo S., Park J., Lee J.-Y., Kweon I. S. CBAM: Convolutional Block Attention Module. ECCV, 2018.",
        "Wang et al. ResU-KAN: U-KAN based residual model for medical image segmentation. 本地课程材料。",
        "Xu et al. KC-UNet: U-Net with Tokenized KAN modules and attention for medical image segmentation. 本地课程材料。",
        "TransUKAN: Computing-Efficient Hybrid KAN-Transformer for Enhanced Medical Image Segmentation. 本地课程材料。",
        "VMKLA-UNet: Vision Mamba with KAN Linear Attention U-Net. 本地课程材料。",
        "A Review of KANs for Medical Image Analysis. 本地课程材料。",
        "Al-Dhabyani W., Gomaa M., Khaled H., Fahmy A. Dataset of Breast Ultrasound Images. Data in Brief, 2020.",
        "Bernal J., Sánchez F. J., Fernández-Esparrach G., et al. WM-DOVA maps for accurate polyp highlighting in colonoscopy. Computerized Medical Imaging and Graphics, 2015.",
    ]
    for i, ref in enumerate(refs, 1):
        add_para(doc, f"[{i}] {ref}", first_line=False)
    doc.save(DOCX_OUT)


def main() -> None:
    figs = create_figures_v2()
    tables = build_table_data()
    stats = main_findings()
    build_markdown_v2(figs, tables, stats)
    build_docx_v2(figs, tables, stats)
    print(f"wrote_markdown={MD_OUT}")
    print(f"wrote_docx={DOCX_OUT}")
    print(f"figures_dir={FIG_DIR}")


if __name__ == "__main__":
    main()
