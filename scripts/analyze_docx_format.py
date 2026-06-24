from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def length_cm(value):
    return round(value.cm, 3) if value is not None else None


def points(value):
    return round(value.pt, 2) if value is not None else None


def run_record(run):
    rpr = run._element.rPr
    rfonts = rpr.rFonts if rpr is not None else None
    east_asia = rfonts.get(qn("w:eastAsia")) if rfonts is not None else None
    ascii_font = rfonts.get(qn("w:ascii")) if rfonts is not None else None
    return {
        "text": run.text[:120],
        "font": run.font.name or ascii_font,
        "east_asia": east_asia,
        "size_pt": points(run.font.size),
        "bold": run.bold,
        "italic": run.italic,
    }


def paragraph_record(index, paragraph):
    fmt = paragraph.paragraph_format
    try:
        alignment = str(paragraph.alignment)
    except ValueError:
        ppr = paragraph._p.pPr
        alignment = (
            ppr.jc.get(qn("w:val"))
            if ppr is not None and ppr.jc is not None
            else None
        )
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text[:300],
        "alignment": alignment,
        "left_indent_cm": length_cm(fmt.left_indent),
        "right_indent_cm": length_cm(fmt.right_indent),
        "first_line_indent_cm": length_cm(fmt.first_line_indent),
        "space_before_pt": points(fmt.space_before),
        "space_after_pt": points(fmt.space_after),
        "line_spacing": (
            round(fmt.line_spacing, 3)
            if isinstance(fmt.line_spacing, float)
            else points(fmt.line_spacing)
        ),
        "line_spacing_rule": str(fmt.line_spacing_rule),
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "runs": [run_record(run) for run in paragraph.runs if run.text][:8],
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    section_records = []
    for index, section in enumerate(document.sections, start=1):
        section_records.append(
            {
                "index": index,
                "width_cm": length_cm(section.page_width),
                "height_cm": length_cm(section.page_height),
                "top_margin_cm": length_cm(section.top_margin),
                "bottom_margin_cm": length_cm(section.bottom_margin),
                "left_margin_cm": length_cm(section.left_margin),
                "right_margin_cm": length_cm(section.right_margin),
                "header_distance_cm": length_cm(section.header_distance),
                "footer_distance_cm": length_cm(section.footer_distance),
                "different_first_page": section.different_first_page_header_footer,
                "start_type": str(section.start_type),
            }
        )

    paragraphs = [
        paragraph_record(index, paragraph)
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    style_counts = Counter(record["style"] for record in paragraphs)
    run_fonts = Counter()
    run_sizes = Counter()
    for record in paragraphs:
        for run in record["runs"]:
            run_fonts[(run["east_asia"], run["font"])] += len(run["text"])
            run_sizes[run["size_pt"]] += len(run["text"])

    result = {
        "file": str(args.docx.resolve()),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_image_count": len(document.inline_shapes),
        "sections": section_records,
        "style_counts": dict(style_counts),
        "run_fonts_by_characters": [
            {"east_asia": key[0], "ascii": key[1], "characters": value}
            for key, value in run_fonts.most_common()
        ],
        "run_sizes_by_characters": [
            {"size_pt": key, "characters": value}
            for key, value in run_sizes.most_common()
        ],
        "paragraphs": paragraphs,
    }
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        args.out.write_text(output, encoding="utf-8")
    else:
        print(output)


if __name__ == "__main__":
    main()
