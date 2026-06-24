from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document


REPLACEMENTS = [
    (
        "利用 tokenized KAN block 加强中高层特征的非线性表达",
        "利用 token 化 KANBlock 增强中高层特征的非线性表达",
    ),
    ("tokenized 中间表征", "token 化的中间表征"),
    ("tokenized 高层表征", "token 化的高层表征"),
    (
        "本文在本地 PyTorch 工程中对 U-KAN 进行课程级复现",
        "本文在本地 PyTorch 实验框架中对 U-KAN 进行了系统复现",
    ),
    (
        "补充实验进一步证明：KAN 的优势不是单一 seed 的偶然峰值",
        "补充实验进一步表明：KAN 的优势并非由单一随机种子下的偶然高值造成",
    ),
    (
        "本文形成了从数据整理、模型复现、训练评估、异常诊断到 GitHub 工程管理的完整闭环。",
        "本文构建了覆盖数据整理、模型复现、训练评估、异常诊断与工程管理的完整研究实现流程。",
    ),
    (
        "若只运行官方代码并复述论文结论，课程实践价值有限",
        "若仅运行官方代码并复述原文结论，则难以充分体现课程实践的实验设计与分析价值",
    ),
    ("适合构造严谨消融", "适合开展较为严格的消融实验设计"),
    (
        "四个问题分别由 U-Net 对照、no-KAN 消融、Attention-U-KAN 对照、随机种子与继续训练补实验回答。",
        "上述四个问题分别通过 U-Net 对照、no-KAN 消融、Attention-U-KAN 对照以及随机种子与继续训练补充实验进行检验。",
    ),
    (
        "本文不使用自绘 U-Net 架构图，而通过上述结构参数、代码路径和最终效率表描述该基线，避免图示质量影响论文表达。",
        "鉴于该基线结构较为成熟，本文主要通过网络配置、代码实现与效率结果对其进行说明。",
    ),
    (
        "KAN 的非线性不只发生在节点上，而是分布在输入—输出连接本身。",
        "KAN 的非线性并非仅由节点激活提供，而是进一步体现在连接边上的可学习一元函数中。",
    ),
    (
        "将 128 通道特征映射为 160 维 token，再将 160 维映射到 256 维 token。",
        "将 128 通道特征映射为 160 维 token 表征，再进一步映射为 256 维 token 表征。",
    ),
    (
        "因此，U-KAN 与 no-KAN 的指标差异更接近对 KANLinear 本身贡献的控制变量分析。",
        "因此，两者的指标差异更适合作为评估 KANLinear 贡献的控制变量结果。",
    ),
    (
        "需要强调的是，该结构是一项待实验检验的改进假设，而不是预设有效的结论。",
        "该结构仅作为改进方案提出，其有效性仍需通过实验结果验证。",
    ),
    (
        "U-Net 回答“U-KAN 相比经典 CNN 是否具有优势”；no-KAN 回答“收益是否来自 KANLinear”；Attention-U-KAN 回答“简单通道—空间注意力是否形成互补”。",
        "U-Net 作为基线，用于检验 U-KAN 相对经典 CNN 的性能变化；no-KAN 用于检验性能收益是否来自 KANLinear；Attention-U-KAN 用于检验通道—空间注意力是否能够与原模型形成互补。",
    ),
    (
        "若继续训练只产生新的末期峰值而未超过原值，就不能继续用“再多跑一些也许会更高”解释差距。",
        "若继续训练仅产生新的末期波动而未超过原始最优值，则难以继续将差距归因于训练轮数不足。",
    ),
    (
        "BUSI no-KAN 的 loss 曲线整体高于其他模型",
        "BUSI no-KAN 的训练损失曲线整体高于其他模型",
    ),
    ("CVC 中 U-Net 的 best epoch 靠近 99", "CVC 中 U-Net 的最佳轮次接近第 99 轮"),
    (
        "首先分析“loss 更高但指标更好”的表面矛盾。",
        "首先分析“训练损失较高而验证指标较优”这一表面矛盾。",
    ),
    (
        "训练 loss 是训练集在数据增强和连续概率下的平均目标",
        "训练损失是训练集在数据增强和连续概率空间下的平均目标",
    ),
    ("loss 更低", "训练损失更低"),
    ("训练 loss 曲线", "训练损失曲线"),
    (
        "其次分析局部“蹦升”是否意味着模型蒙对。",
        "其次，需要分析局部指标突升是否仅由随机因素造成。",
    ),
    (
        "这说明 KAN 的收益不是 seed2981 的单次幸运峰值",
        "这说明 KAN 的收益并非由 seed2981 下的单次随机高值所致",
    ),
    (
        "U-KAN 的精度提升不是免费收益",
        "U-KAN 的性能提升以额外计算开销为代价",
    ),
    (
        "工程证据的作用不是装饰论文，而是建立结果的可追溯链",
        "这些工程记录并非附属性材料，而是结果可追溯性的重要组成部分",
    ),
    (
        "不能把所有表格压缩为单一“冠军模型”",
        "不能将全部结果简单归结为单一最优模型",
    ),
    (
        "不能推导为 KAN 在所有分割任务上普遍优于 CNN 或 MLP",
        "不能据此推断 KAN 在所有分割任务中都普遍优于 CNN 或 MLP",
    ),
    (
        "不宣称完全复刻官方数值",
        "不将本地结果表述为对官方数值的严格再现",
    ),
    (
        "直觉上这可能像“偶然蒙对”，但直觉本身不是证据。",
        "直观上，这一现象可能表现为随机波动导致的偶然高值，但该判断仍需实验验证。",
    ),
    (
        "最终把结论从“某次最优值更高”修正为“不同随机条件下相对方向一致”。",
        "因此，结论由“单次最优值更高”调整为“在不同随机条件下保持相同的相对排序方向”。",
    ),
    (
        "本文围绕 U-KAN 医学图像分割算法完成了课程级完整复现。",
        "本文围绕 U-KAN 医学图像分割算法完成了系统复现。",
    ),
    (
        "Attention-U-KAN 没有稳定超过原始 U-KAN，说明简单注意力并非必然有效。",
        "在当前统一训练协议下，Attention-U-KAN 未稳定超过原始 U-KAN，说明简单注意力的引入并不必然带来性能增益。",
    ),
    (
        "最终，项目形成了从数据处理、模型复现、实验诊断、结果讨论到 Git 版本管理和 Word 论文生成的完整闭环",
        "最终，项目建立了覆盖数据处理、模型复现、实验诊断、结果讨论、Git 版本管理与 Word 论文生成的完整工作流程",
    ),
    (
        "提交论文前应再次执行 git push origin course-paper-ukan，确认远程仓库包含最终 Word、README 和实验结果。",
        "为保证仓库内容与论文一致，建议在提交前同步更新远程分支，并确认其中包含最终 Word、README 和实验结果。",
    ),
    (
        "本项目没有自行采集或标注原始临床数据，因此不将 BUSI/CVC 宣称为原创数据集。",
        "本项目没有自行采集或标注原始临床数据，因此本文不将 BUSI 与 CVC-ClinicDB 归类为自建原创数据集。",
    ),
    (
        "本文创新结构保留作者自绘图。",
        "本文提出的改进结构示意图均由作者自行绘制。",
    ),
    ("one-batch smoke test", "单批次冒烟测试"),
    ("smoke test", "冒烟测试"),
    ("per-image Dice", "逐图 Dice"),
    ("现有 per-image 与病灶面积分层结果", "现有逐图统计与病灶面积分层结果"),
    ("单通道二分类 mask logit", "单通道二分类分割 logit"),
    ("KAN layer 用于", "KAN 层（KAN layer）用于"),
    ("KAN/tokenized", "KAN/token 化"),
    ("tokenized U 型", "token 化 U 型"),
    ("tokenized 拓扑", "token 化拓扑"),
    ("tokenized 编码", "token 化编码"),
    ("tokenized 结构", "token 化结构"),
    ("skip fusion", "跳跃连接融合"),
    ("CVC best checkpoint", "CVC 最佳检查点"),
    ("best checkpoint", "最佳检查点"),
    ("best epoch", "最佳轮次"),
    (
        "U-KAN 将 Kolmogorov-Arnold Network 的可学习边函数引入 U 型网络的 token 化的中间表征。",
        "U-KAN 在 U 型网络的 token 化中间表征中引入 Kolmogorov-Arnold Network 的可学习边函数。",
    ),
    ("四个 跳跃连接融合 位置", "四个跳跃连接融合位置"),
    ("在 跳跃连接融合 后", "在跳跃连接融合后"),
    ("四个 跳跃连接融合 后", "四处跳跃连接融合后"),
    ("的 跳跃连接融合", "的跳跃连接融合"),
    ("四处 跳跃连接融合", "四处跳跃连接融合"),
    ("KAN/token 化 模块", "KAN 与 token 化模块"),
    ("KAN/token 化 结构", "KAN 与 token 化结构"),
    ("从 最佳检查点", "从最佳检查点"),
    ("最佳检查点 保存", "最佳检查点保存"),
    ("最佳检查点 继续训练", "最佳检查点继续训练"),
    ("最佳检查点 续训", "最佳检查点续训"),
    ("最佳轮次 靠后", "最佳轮次靠后"),
    ("BCE+Dice loss", "BCE+Dice 损失"),
    ("soft Dice loss", "soft Dice 损失"),
    ("Dice loss", "Dice 损失"),
    ("no-KAN loss 偏高", "no-KAN 的训练损失偏高"),
    ("no-KAN loss 与 IoU", "no-KAN 训练损失与 IoU"),
    ("loss 与阈值化指标", "损失函数与阈值化指标"),
    ("训练 loss", "训练损失"),
    ("从最佳检查点 延续", "从最佳检查点延续"),
    ("最佳轮次 接近", "最佳轮次接近"),
    ("原始 最佳轮次 为", "原始最佳轮次为"),
    ("模型的 最佳轮次 为", "模型的最佳轮次为"),
    (
        "Attention-U-KAN 未稳定超过原始 U-KAN，表明简单通道—空间注意力插入并非必然有效。",
        "在当前统一训练协议下，Attention-U-KAN 未稳定超过原始 U-KAN，表明简单通道—空间注意力的引入并不必然带来性能增益。",
    ),
]


NEEDLES = [
    "课程级",
    "完整闭环",
    "若只运行",
    "严谨消融",
    "回答“",
    "分别由",
    "不使用自绘",
    "非线性不只",
    "160 维 token",
    "更接近对 KANLinear",
    "待实验检验",
    "再多跑",
    "best epoch",
    "loss 更高",
    "蹦升",
    "蒙对",
    "幸运峰值",
    "免费收益",
    "装饰论文",
    "冠军模型",
    "不能推导为",
    "不宣称完全复刻",
    "git push",
    "宣称为原创",
    "保留作者自绘",
    "tokenized",
    "skip fusion",
    "smoke test",
    "per-image",
]


def iter_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_in_paragraph(paragraph, old: str, new: str) -> int:
    count = paragraph.text.count(old)
    for _ in range(count):
        full_text = "".join(run.text for run in paragraph.runs)
        start = full_text.rfind(old)
        if start < 0:
            break
        end = start + len(old)

        positions = []
        cursor = 0
        for index, run in enumerate(paragraph.runs):
            run_start = cursor
            run_end = cursor + len(run.text)
            positions.append((index, run_start, run_end))
            cursor = run_end

        first = next(
            index for index, run_start, run_end in positions if run_end > start
        )
        last = next(
            index for index, run_start, run_end in positions if run_end >= end
        )
        first_start = positions[first][1]
        last_start = positions[last][1]
        prefix = paragraph.runs[first].text[: start - first_start]
        suffix = paragraph.runs[last].text[end - last_start :]
        paragraph.runs[first].text = prefix + new + suffix
        for index in range(first + 1, last + 1):
            paragraph.runs[index].text = ""
    return count


def apply_revisions(path: Path, backup: Path | None) -> None:
    if backup is not None and not backup.exists():
        shutil.copy2(path, backup)

    document = Document(path)
    total = 0
    for old, new in REPLACEMENTS:
        matches = 0
        for paragraph in iter_paragraphs(document):
            matches += replace_in_paragraph(paragraph, old, new)
        if matches:
            print(f"replaced={matches}: {old} -> {new}")
            total += matches

    document.save(path)
    print(f"total_replacements={total}")
    print(f"saved={path.resolve()}")
    if backup is not None:
        print(f"backup={backup.resolve()}")


def inspect(path: Path) -> None:
    document = Document(path)
    for index, paragraph in enumerate(iter_paragraphs(document)):
        if any(needle in paragraph.text for needle in NEEDLES):
            print(f"{index}: {paragraph.text}")


def compare(before: Path, after: Path) -> None:
    old = Document(before)
    new = Document(after)
    old_paragraphs = list(iter_paragraphs(old))
    new_paragraphs = list(iter_paragraphs(new))
    print(f"paragraphs={len(old_paragraphs)}->{len(new_paragraphs)}")
    print(f"tables={len(old.tables)}->{len(new.tables)}")
    print(f"images={len(old.inline_shapes)}->{len(new.inline_shapes)}")
    print(f"sections={len(old.sections)}->{len(new.sections)}")
    assert len(old_paragraphs) == len(new_paragraphs)
    assert len(old.tables) == len(new.tables)
    assert len(old.inline_shapes) == len(new.inline_shapes)
    assert len(old.sections) == len(new.sections)

    changes = [
        (index, old_p.text, new_p.text)
        for index, (old_p, new_p) in enumerate(zip(old_paragraphs, new_paragraphs))
        if old_p.text != new_p.text
    ]
    print(f"changed_paragraphs={len(changes)}")
    for index, old_text, new_text in changes:
        print(f"CHANGE {index}\n- {old_text}\n+ {new_text}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--backup", type=Path)
    parser.add_argument("--compare-to", type=Path)
    args = parser.parse_args()
    if args.compare_to:
        compare(args.compare_to, args.docx)
    elif args.inspect:
        inspect(args.docx)
    else:
        apply_revisions(args.docx, args.backup)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
