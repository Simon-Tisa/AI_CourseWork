from __future__ import annotations

import csv
import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps


ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parent
PAPER_DIR = ROOT / "paper"
FIG_DIR = PAPER_DIR / "figures"
EXP_FIG_DIR = ROOT / "experiments" / "figures"
TABLE_DIR = PAPER_DIR / "tables"
COURSE_IMG_DIR = WORKSPACE / "课程论文要求" / "图片"
HANDDRAWN_PDF = WORKSPACE / "fig-PPT手绘.pdf"

TITLE = "基于 U-KAN 的医学图像分割算法复现与注意力增强改进研究"
DOCX_OUT = PAPER_DIR / "基于U-KAN的医学图像分割算法复现与注意力增强改进研究.docx"
MD_OUT = PAPER_DIR / "course_paper_draft.md"
GITHUB_URL = "https://github.com/Simon-Tisa/AI_CourseWork.git"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def fmt4(value: str | float) -> str:
    return f"{float(value):.4f}"


def fmt2(value: str | float) -> str:
    return f"{float(value):.2f}"


def fmt_params(value: str | float) -> str:
    return f"{float(value) / 1_000_000:.2f}M"


def find_row(rows: list[dict[str, str]], name: str) -> dict[str, str]:
    for row in rows:
        if row["name"] == name:
            return row
    raise KeyError(name)


def rel(path: Path) -> str:
    return path.relative_to(PAPER_DIR).as_posix()


def font_path(bold: bool = False) -> str:
    candidates = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for item in candidates:
        if Path(item).exists():
            return item
    return candidates[-1]


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path(bold), size=size)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> tuple[int, int]:
    if not text:
        return 0, 0
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for char in paragraph:
            trial = current + char
            if text_size(draw, trial, font)[0] <= max_width or not current:
                current = trial
            else:
                lines.append(current)
                current = char
        lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.ImageFont,
    fill: str,
    max_width: int,
    line_gap: int = 8,
    align: str = "left",
) -> int:
    x, y = xy
    lines = wrap_text(draw, text, font, max_width)
    line_heights = [text_size(draw, line, font)[1] for line in lines]
    for line, h in zip(lines, line_heights):
        w, _ = text_size(draw, line, font)
        dx = {"left": 0, "center": (max_width - w) // 2, "right": max_width - w}.get(align, 0)
        draw.text((x + dx, y), line, font=font, fill=fill)
        y += h + line_gap
    return y


def draw_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    fill: str,
    outline: str,
    font: ImageFont.ImageFont,
    text_fill: str = "#17212b",
    radius: int = 22,
    width: int = 3,
    align: str = "center",
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    max_width = x2 - x1 - 32
    lines = wrap_text(draw, text, font, max_width)
    heights = [text_size(draw, line, font)[1] for line in lines]
    total_h = sum(heights) + (len(lines) - 1) * 7
    y = y1 + max(0, (y2 - y1 - total_h) // 2)
    for line, h in zip(lines, heights):
        w, _ = text_size(draw, line, font)
        if align == "center":
            x = x1 + (x2 - x1 - w) // 2
        else:
            x = x1 + 18
        draw.text((x, y), line, font=font, fill=text_fill)
        y += h + 7


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str, width: int = 5) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    p1 = (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6))
    p2 = (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, p1, p2], fill=fill)


def add_title(draw: ImageDraw.ImageDraw, text: str, width: int, y: int) -> None:
    font = pil_font(48, True)
    w, h = text_size(draw, text, font)
    draw.text(((width - w) // 2, y), text, font=font, fill="#16202a")


def create_kan_vs_mlp() -> Path:
    out = FIG_DIR / "kan_vs_mlp_diagram.png"
    W, H = 2000, 1120
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(40, True)
    label_font = pil_font(29, True)
    body_font = pil_font(25)
    small_font = pil_font(21)
    add_title(draw, "MLP 与 KAN 的非线性建模差异", W, 44)

    draw.rounded_rectangle((70, 145, 955, 965), radius=28, fill="#ffffff", outline="#cbd5e1", width=3)
    draw.rounded_rectangle((1045, 145, 1930, 965), radius=28, fill="#ffffff", outline="#cbd5e1", width=3)
    draw.text((110, 180), "MLP：节点固定激活", font=title_font, fill="#2563eb")
    draw.text((1085, 180), "KAN：边上可学习函数", font=title_font, fill="#047857")

    def draw_network(x0: int, y0: int, color: str, kan: bool) -> None:
        layer_x = [x0, x0 + 250, x0 + 500]
        ys = [[y0 + 115, y0 + 280, y0 + 445], [y0 + 80, y0 + 245, y0 + 410, y0 + 575], [y0 + 160, y0 + 325, y0 + 490]]
        for i, xs in enumerate(layer_x[:-1]):
            for y1 in ys[i]:
                for y2 in ys[i + 1]:
                    line_color = "#9ca3af" if not kan else color
                    draw.line([(xs + 32, y1), (layer_x[i + 1] - 32, y2)], fill=line_color, width=2 if not kan else 3)
                    if kan and (int(y1 + y2) % 3 == 0):
                        mx = (xs + layer_x[i + 1]) // 2
                        my = (y1 + y2) // 2 - 8
                        draw.text((mx - 26, my), "φ(x)", font=small_font, fill="#047857")
        for lx, ly in zip(layer_x, ys):
            for y in ly:
                draw.ellipse((lx - 32, y - 32, lx + 32, y + 32), fill="#f8fafc", outline=color, width=5)
        if not kan:
            for lx, ly in zip(layer_x[1:], ys[1:]):
                for y in ly:
                    draw.rounded_rectangle((lx - 54, y + 43, lx + 54, y + 88), radius=12, fill="#dbeafe", outline="#93c5fd")
                    draw.text((lx - 47, y + 52), "ReLU", font=small_font, fill="#1d4ed8")

    draw_network(210, 300, "#2563eb", False)
    draw_network(1185, 300, "#047857", True)

    draw_box(
        draw,
        (125, 785, 900, 920),
        "边上只有可学习标量权重，节点使用 ReLU/GELU 等固定激活函数；非线性主要发生在节点处。",
        "#eff6ff",
        "#bfdbfe",
        body_font,
        align="left",
    )
    draw_box(
        draw,
        (1100, 785, 1875, 920),
        "每条边学习一维函数 φe(x)，通常用 B-spline 基函数展开：φe(x)=Σ ci Bi(x)；节点主要对入边输出求和。",
        "#ecfdf5",
        "#a7f3d0",
        body_font,
        align="left",
    )
    draw.text(
        (90, 1015),
        f"图源说明：作者根据手绘草图 {HANDDRAWN_PDF.name} 与 KAN 论文思想重新绘制。",
        font=small_font,
        fill="#475569",
    )
    img.save(out, quality=95)
    return out


def create_ukan_architecture() -> Path:
    out = FIG_DIR / "ukan_architecture_diagram.png"
    W, H = 2550, 1200
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "本文复现的 U-KAN 医学图像分割结构", W, 42)
    body = pil_font(23)
    small = pil_font(20)

    boxes = [
        ("Input\n256×256×3", 90, 510, 190, "#e0f2fe", "#0284c7"),
        ("Conv\n16ch", 350, 440, 190, "#eef2ff", "#4f46e5"),
        ("Conv\n32ch", 610, 370, 190, "#eef2ff", "#4f46e5"),
        ("Conv\n128ch", 870, 300, 200, "#eef2ff", "#4f46e5"),
        ("Patch + KANBlock\n160ch", 1140, 230, 250, "#ecfdf5", "#059669"),
        ("Patch + KANBlock\n256ch", 1470, 300, 250, "#ecfdf5", "#059669"),
        ("Up + KANBlock\n160ch", 1800, 370, 240, "#fef3c7", "#d97706"),
        ("Up + KANBlock\n128ch", 2040, 440, 240, "#fef3c7", "#d97706"),
        ("Up Conv\n32/16ch", 2040, 590, 240, "#fff7ed", "#ea580c"),
        ("Output Mask\n1ch", 2320, 510, 190, "#fce7f3", "#db2777"),
    ]
    skips = [
        ((445, 440), (2160, 590), "skip1"),
        ((705, 370), (2160, 590), "skip2"),
        ((970, 300), (2160, 440), "skip3"),
        ((1265, 230), (1920, 370), "skip4"),
    ]
    for start, end, label in skips:
        mid_y = min(start[1], end[1]) - 85
        draw.line([start, (start[0], mid_y), (end[0], mid_y), end], fill="#94a3b8", width=4)
        if label in {"skip3", "skip4"}:
            draw.text(((start[0] + end[0]) // 2 - 35, mid_y - 30), label, font=small, fill="#475569")

    centers: dict[str, tuple[int, int]] = {}
    for label, x, y, box_w, fill, outline in boxes:
        draw_box(draw, (x, y, x + box_w, y + 105), label, fill, outline, body)
        centers[label] = (x + box_w // 2, y + 52)
    order = boxes[:8] + [boxes[8], boxes[9]]
    for left, right in zip(order, order[1:]):
        start = (left[1] + left[3], left[2] + 52)
        end = (right[1], right[2] + 52)
        arrow(draw, start, end, "#64748b", 4)

    draw_box(
        draw,
        (270, 820, 940, 1035),
        "复现重点：低层卷积提取局部纹理；中高层转为 token，并用 KANBlock 替代 MLP 非线性映射。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw_box(
        draw,
        (1010, 820, 1700, 1035),
        "消融模型 U-KAN(no-KAN)：保持同一 U 型结构和训练设置，仅把 KANLinear 替换为普通 Linear，用于隔离 KAN 模块贡献。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw.text((85, 1090), "图源说明：作者根据 U-KAN 官方源码结构和本文 PyTorch 复现代码整理绘制。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def create_attention_module() -> Path:
    out = FIG_DIR / "attention_ukan_module.png"
    W, H = 1900, 1050
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "Attention-U-KAN 的 skip fusion 增强模块", W, 42)
    body = pil_font(24)
    small = pil_font(20)
    draw_box(draw, (120, 340, 360, 455), "Decoder\nfeature", "#eef2ff", "#4f46e5", body)
    draw_box(draw, (120, 540, 360, 655), "Encoder\nskip feature", "#e0f2fe", "#0284c7", body)
    draw_box(draw, (520, 440, 730, 555), "Add\n融合", "#f8fafc", "#64748b", body)
    draw_box(draw, (875, 330, 1255, 470), "Channel attention\nGAP → 1×1 Conv\n→ Sigmoid", "#ecfdf5", "#059669", body)
    draw_box(draw, (875, 560, 1255, 700), "Spatial attention\nAvg/Max → 7×7 Conv\n→ Sigmoid", "#fff7ed", "#ea580c", body)
    draw_box(draw, (1390, 440, 1630, 555), "Refined\nfeature", "#fce7f3", "#db2777", body)
    draw_box(draw, (1690, 440, 1840, 555), "Next\nblock", "#f8fafc", "#64748b", body)

    arrow(draw, (360, 398), (520, 485), "#64748b", 5)
    arrow(draw, (360, 598), (520, 510), "#64748b", 5)
    arrow(draw, (730, 498), (875, 400), "#64748b", 5)
    arrow(draw, (1065, 470), (1065, 560), "#64748b", 5)
    arrow(draw, (1255, 630), (1390, 498), "#64748b", 5)
    arrow(draw, (1630, 498), (1690, 498), "#64748b", 5)

    draw_box(
        draw,
        (270, 780, 1620, 925),
        "设计动机：skip connection 直接相加可能把背景纹理和边界噪声带入解码器；通道-空间注意力用于重标定融合后的特征。但实验显示该简单插入方式未稳定优于原始 U-KAN。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw.text((90, 980), "图源说明：作者根据本文 attention.py 与 attention_ukan.py 代码自绘。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def create_workflow() -> Path:
    out = FIG_DIR / "experiment_workflow.png"
    W, H = 2100, 1100
    img = Image.new("RGB", (W, H), "#fbfcfd")
    draw = ImageDraw.Draw(img)
    add_title(draw, "实验流程与课程论文产出链路", W, 42)
    body = pil_font(24)
    small = pil_font(20)
    items = [
        ("公开数据集\nBUSI / CVC", "#e0f2fe", "#0284c7"),
        ("数据整理\nmask 合并、二值化\n固定 train/val split", "#eef2ff", "#4f46e5"),
        ("四组模型\nU-Net / no-KAN\nU-KAN / Attention-U-KAN", "#ecfdf5", "#059669"),
        ("统一训练\n100 epoch\nBCE+Dice Loss", "#fef3c7", "#d97706"),
        ("统一评估\nIoU/Dice/Precision\nRecall/Specificity/耗时", "#fff7ed", "#ea580c"),
        ("补充实验\nBUSI seed6142\nCVC 继续训练", "#fce7f3", "#db2777"),
        ("论文与 GitHub\n图表、README\n复现命令、结论", "#f8fafc", "#64748b"),
    ]
    x = 90
    y = 355
    prev_end = None
    for idx, (text, fill, outline) in enumerate(items):
        w = 250 if idx not in [1, 2, 4] else 310
        draw_box(draw, (x, y, x + w, y + 160), text, fill, outline, body)
        if prev_end:
            arrow(draw, prev_end, (x, y + 80), "#64748b", 5)
        prev_end = (x + w, y + 80)
        x += w + 50
    draw_box(
        draw,
        (160, 720, 1940, 910),
        "关键控制：所有主实验使用同一输入分辨率 256、batch size 8、seed 2981、Adam 与 CosineAnnealingLR；最终量化结果统一来自 evaluate.py 的全验证集像素级混淆矩阵口径。",
        "#ffffff",
        "#cbd5e1",
        body,
        align="left",
    )
    draw.text((90, 1000), "图源说明：作者根据本文实验脚本、数据处理脚本和论文写作流程自绘。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def create_screenshot_collage() -> Path:
    out = FIG_DIR / "runtime_screenshots_collage.png"
    W, H = 1900, 1200
    img = Image.new("RGB", (W, H), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(42, True)
    body = pil_font(24)
    small = pil_font(18)
    title = "训练、评估与结果生成过程截图"
    tw, _ = text_size(draw, title, title_font)
    draw.text(((W - tw) // 2, 38), title, font=title_font, fill="#17212b")
    sources = [
        ("BUSI U-KAN 训练完成", COURSE_IMG_DIR / "busi_ukan代码最终过程.png"),
        ("BUSI Attention-U-KAN 训练完成", COURSE_IMG_DIR / "busi_attention_ukan最终代码运行过程.png"),
        ("CVC U-KAN 训练完成", COURSE_IMG_DIR / "cvc_ukan代码运行最终过程-epoch100.png"),
        ("metrics.csv 与预测 mask 生成", COURSE_IMG_DIR / "mess" / "生成 metrics.csv 和预测 mask.png"),
    ]
    cells = [(70, 150), (985, 150), (70, 665), (985, 665)]
    box_w, box_h = 845, 420
    for (label, path), (x, y) in zip(sources, cells):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=20, fill="#ffffff", outline="#cbd5e1", width=3)
        draw.text((x + 20, y + 18), label, font=body, fill="#0f172a")
        if path.exists():
            src = Image.open(path).convert("RGB")
            src.thumbnail((box_w - 40, box_h - 80), Image.Resampling.LANCZOS)
            px = x + (box_w - src.width) // 2
            py = y + 65 + (box_h - 85 - src.height) // 2
            draw.rounded_rectangle((px - 5, py - 5, px + src.width + 5, py + src.height + 5), radius=8, fill="#e2e8f0")
            img.paste(src, (px, py))
        else:
            draw.text((x + 40, y + 180), f"未找到截图：{path.name}", font=body, fill="#b91c1c")
    draw.text((70, 1110), "图源说明：作者在本地 Windows/Anaconda/PyTorch 环境运行项目时截取，本文整理为拼图。", font=small, fill="#475569")
    img.save(out, quality=95)
    return out


def copy_course_reference_files() -> None:
    if HANDDRAWN_PDF.exists():
        shutil.copy2(HANDDRAWN_PDF, PAPER_DIR / HANDDRAWN_PDF.name)


def generate_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    copy_course_reference_files()
    figures = {
        "kan_vs_mlp": create_kan_vs_mlp(),
        "ukan_architecture": create_ukan_architecture(),
        "attention_module": create_attention_module(),
        "workflow": create_workflow(),
        "screenshots": create_screenshot_collage(),
        "busi_samples": EXP_FIG_DIR / "busi_samples.png",
        "cvc_samples": EXP_FIG_DIR / "cvc_samples.png",
        "busi_curves": EXP_FIG_DIR / "busi_training_curves.png",
        "cvc_curves": EXP_FIG_DIR / "cvc_training_curves.png",
        "busi_preds": EXP_FIG_DIR / "busi_prediction_comparison.png",
        "cvc_preds": EXP_FIG_DIR / "cvc_prediction_comparison.png",
    }
    return figures


def markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    out = ["| " + " | ".join(headers) + " |"]
    out.append("| " + " | ".join(["---"] * len(headers)) + " |")
    out.extend("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join(out)


def build_table_data() -> dict[str, list[list[str]]]:
    results = read_csv(TABLE_DIR / "segmentation_results.csv")
    diagnostics = read_csv(TABLE_DIR / "training_diagnostics.csv")
    busi_stats = read_csv(ROOT / "experiments" / "results" / "busi_dataset_stats.csv")[0]
    cvc_stats = read_csv(ROOT / "experiments" / "results" / "cvc_dataset_stats.csv")[0]
    seed_stability = read_csv(TABLE_DIR / "busi_seed_stability.csv")
    cvc_continue = read_csv(TABLE_DIR / "cvc_continued_training.csv")

    def result_rows(dataset: str) -> list[list[str]]:
        names = [
            f"{dataset}_unet_seed2981",
            f"{dataset}_no_kan_seed2981",
            f"{dataset}_ukan_seed2981",
            f"{dataset}_attention_ukan_seed2981",
        ]
        rows: list[list[str]] = []
        for name in names:
            r = find_row(results, name)
            rows.append(
                [
                    r["variant"],
                    fmt4(r["iou"]),
                    fmt4(r["dice"]),
                    fmt4(r["precision"]),
                    fmt4(r["recall"]),
                    fmt4(r["specificity"]),
                    fmt_params(r["params"]),
                    fmt2(r["infer_ms_per_image"]),
                ]
            )
        return rows

    diag_names = [
        "busi_unet_seed2981",
        "busi_no_kan_seed2981",
        "busi_ukan_seed2981",
        "busi_attention_ukan_seed2981",
        "cvc_unet_seed2981",
        "cvc_no_kan_seed2981",
        "cvc_ukan_seed2981",
        "cvc_attention_ukan_seed2981",
    ]
    diag_rows = []
    for name in diag_names:
        r = find_row(diagnostics, name)
        diag_rows.append(
            [
                name.replace("_seed2981", ""),
                r["best_epoch"],
                fmt4(r["log_best_val_iou"]),
                fmt4(r["final_val_iou"]),
                fmt4(r["last10_val_iou_delta"]),
                fmt4(r["evaluate_iou"]),
            ]
        )

    return {
        "dataset": [
            [
                "BUSI",
                busi_stats["num_images"],
                f"{busi_stats['min_width']}–{busi_stats['max_width']}",
                f"{busi_stats['min_height']}–{busi_stats['max_height']}",
                fmt4(busi_stats["mean_mask_ratio"]),
                "乳腺超声病灶分割；多 mask 病例已合并",
            ],
            [
                "CVC-ClinicDB",
                cvc_stats["num_images"],
                f"{cvc_stats['min_width']}–{cvc_stats['max_width']}",
                f"{cvc_stats['min_height']}–{cvc_stats['max_height']}",
                fmt4(cvc_stats["mean_mask_ratio"]),
                "结肠镜息肉分割；图像与 mask 一一对应",
            ],
        ],
        "platform": [
            ["Python", "3.10.18", "Anaconda torch 环境"],
            ["PyTorch", "2.9.1+cu126", "CUDA 12.6"],
            ["GPU", "NVIDIA GeForce RTX 4080 Laptop GPU", "单卡训练与评估"],
            ["输入分辨率", "256×256", "所有模型统一"],
            ["Batch size / epoch", "8 / 100", "主实验统一设置"],
            ["优化器 / 学习率", "Adam；lr=1e-4，kan_lr=1e-2", "CosineAnnealingLR，min_lr=1e-5"],
            ["损失函数", "0.5×BCEWithLogits + DiceLoss", "二分类 mask 分割"],
            ["随机种子", "2981；补充 6142", "固定划分并记录 split 文件"],
        ],
        "busi_results": result_rows("busi"),
        "cvc_results": result_rows("cvc"),
        "diagnostics": diag_rows,
        "seed_stability": [
            [
                row["variant"],
                fmt4(row["seed2981_iou"]),
                fmt4(row["seed2981_dice"]),
                fmt4(row["seed6142_iou"]),
                fmt4(row["seed6142_dice"]),
                fmt4(row["mean_iou"]),
                fmt4(row["mean_dice"]),
            ]
            for row in seed_stability
        ],
        "cvc_continue": [
            [
                row["name"],
                row["setting"],
                row["epochs"],
                row["best_epoch"],
                fmt4(row["iou"]),
                fmt4(row["dice"]),
                fmt4(row["last10_val_iou_delta"]),
            ]
            for row in cvc_continue
        ],
    }


def main_findings() -> dict[str, str]:
    rows = read_csv(TABLE_DIR / "segmentation_results.csv")
    busi_ukan = find_row(rows, "busi_ukan_seed2981")
    busi_nokan = find_row(rows, "busi_no_kan_seed2981")
    busi_unet = find_row(rows, "busi_unet_seed2981")
    busi_att = find_row(rows, "busi_attention_ukan_seed2981")
    cvc_ukan = find_row(rows, "cvc_ukan_seed2981")
    cvc_nokan = find_row(rows, "cvc_no_kan_seed2981")
    cvc_unet = find_row(rows, "cvc_unet_seed2981")
    cvc_att = find_row(rows, "cvc_attention_ukan_seed2981")
    ft = find_row(rows, "cvc_ukan_seed2981_ft50")
    return {
        "busi_ukan_iou": fmt4(busi_ukan["iou"]),
        "busi_ukan_dice": fmt4(busi_ukan["dice"]),
        "busi_gain_nokan_iou": fmt4(float(busi_ukan["iou"]) - float(busi_nokan["iou"])),
        "busi_gain_unet_iou": fmt4(float(busi_ukan["iou"]) - float(busi_unet["iou"])),
        "busi_att_gap": fmt4(float(busi_ukan["iou"]) - float(busi_att["iou"])),
        "cvc_ukan_iou": fmt4(cvc_ukan["iou"]),
        "cvc_ukan_dice": fmt4(cvc_ukan["dice"]),
        "cvc_gain_nokan_iou": fmt4(float(cvc_ukan["iou"]) - float(cvc_nokan["iou"])),
        "cvc_gain_unet_iou": fmt4(float(cvc_ukan["iou"]) - float(cvc_unet["iou"])),
        "cvc_att_gap": fmt4(float(cvc_ukan["iou"]) - float(cvc_att["iou"])),
        "cvc_ft_gap": fmt4(float(cvc_ukan["iou"]) - float(ft["iou"])),
    }


def build_markdown(figs: dict[str, Path], tables: dict[str, list[list[str]]], stats: dict[str, str]) -> None:
    result_headers = ["模型", "IoU", "Dice", "Precision", "Recall", "Specificity", "参数量", "推理 ms/图"]
    md = f"""# {TITLE}

**课程方向**：方向（一）复现任意人工智能算法  
**项目 GitHub**：[{GITHUB_URL}]({GITHUB_URL})  
**实验代码分支**：`course-paper-ukan`  
**作者**：蔡雪峰  
**日期**：2026 年 6 月

## 来源说明

本文围绕 U-KAN 在二维医学图像分割任务中的复现与改进展开。算法来源主要包括 U-KAN 论文与官方开源项目、Kolmogorov-Arnold Network（KAN）论文及 pykan 项目、U-Net 经典医学图像分割框架、通道-空间注意力机制相关论文，以及 BUSI 与 CVC-ClinicDB 两个公开医学图像分割数据集。项目代码为课程论文重新组织的 PyTorch 工程，未直接复制官方训练脚本；官方源码主要用于核对模型结构、训练超参数和结果解释。

本项目使用的全部实验指标来自本地训练和 `evaluate.py` 脚本评估，未由大语言模型生成或改写。本文写作、结构组织、语言润色和图表说明使用 OpenAI Codex 辅助完成；实验方案、训练数据、运行截图、结果表格和结论均基于本地工程记录。原始手绘草图文件为 `{HANDDRAWN_PDF.name}`，本文进一步绘制了规范化 KAN 原理图、U-KAN 结构图、Attention-U-KAN 模块图和实验流程图，并在每幅图下标明来源。

## 摘要

医学图像分割需要在有限标注数据条件下准确定位病灶或器官边界。传统 U-Net 结构具有清晰的编码器-解码器和跳跃连接，但中高层非线性表达通常依赖卷积或普通 MLP。U-KAN 将 Kolmogorov-Arnold Network 中“边上可学习一维函数”的思想引入 U 型分割网络，在 token 化特征上使用 KANBlock 建模复杂非线性关系。本文选择课程方向（一）“复现任意人工智能算法”，在本地 PyTorch 环境中复现 U-KAN 的医学图像分割流程，并在此基础上尝试加入轻量通道-空间注意力模块，构建 Attention-U-KAN。

实验使用 BUSI 乳腺超声数据集和 CVC-ClinicDB 结肠镜息肉数据集，统一进行数据整理、固定随机划分、训练、评估和可视化。对比模型包括传统 U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN。结果显示，U-KAN 在 BUSI 上取得 IoU={stats["busi_ukan_iou"]}、Dice={stats["busi_ukan_dice"]}，相对 no-KAN 提升 IoU {stats["busi_gain_nokan_iou"]}，相对 U-Net 提升 IoU {stats["busi_gain_unet_iou"]}；在 CVC 上取得 IoU={stats["cvc_ukan_iou"]}、Dice={stats["cvc_ukan_dice"]}，也为主实验中最佳结果。补充随机种子实验表明，BUSI 上 KAN 模块收益不是单一 seed 的偶然现象；CVC 继续训练实验显示，当前 U-KAN 主模型在 100 epoch 后基本达到平台期。Attention-U-KAN 未稳定超过原始 U-KAN，说明简单注意力插入并不必然提升 KAN 分割网络。本文最终形成完整代码仓库、README、数据整理脚本、训练日志、结果表格、预测可视化和课程论文成稿。

**关键词**：医学图像分割；U-KAN；Kolmogorov-Arnold Network；U-Net；注意力机制；算法复现

## 1 引言

医学图像分割是计算机辅助诊断中的基础任务之一。对于乳腺超声、结肠镜等图像，准确分割病灶区域可以辅助医生估计病灶大小、形态和边界，为后续诊断和随访提供量化依据。与自然图像相比，医学图像常具有样本数量有限、标注成本高、边界模糊、目标面积小和图像噪声强等特点，因此模型不仅需要较强的局部纹理提取能力，也需要在中高层语义上表达复杂非线性关系。

U-Net 是医学图像分割中最经典的卷积网络结构。它通过编码器提取多尺度语义特征，通过解码器逐步恢复空间分辨率，并利用跳跃连接弥补上采样过程中的细节损失。此类结构简单、稳定、易训练，但在复杂组织边界和弱对比目标上仍可能受到卷积感受野和普通非线性变换的限制。KAN 提供了另一种建模思路：传统 MLP 通常在节点上放置固定激活函数，而 KAN 将可学习的一维函数放在网络边上，并用样条基函数进行参数化。这种表示方式为神经网络的非线性建模提供了新的可解释结构。

U-KAN 的核心动机是把 KAN 的非线性表达能力嵌入 U 型医学图像分割架构。它保留 U-Net 的局部卷积编码和跳跃连接，同时在更高层的 token 化特征上使用 KANBlock 替代普通 MLP，以增强全局语义和非线性映射能力。本文选择 U-KAN 作为复现对象，是因为它既符合“复现一个人工智能算法”的课程要求，又能自然扩展出医学图像分割数据整理、模型消融、对比实验、可视化和工程化复现等加分项。

本文的主要工作包括：第一，重新组织 U-KAN 分割工程，完成 BUSI 与 CVC 数据集预处理、固定划分、训练和评估；第二，实现 U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN 四类模型，形成主实验矩阵；第三，围绕训练曲线异常和 CVC 收敛性疑问设计补充实验；第四，生成完整图表、运行截图、README 和论文正文，从实验事实出发讨论复现效果与局限。

## 2 算法原理与模型构建

### 2.1 KAN 与 MLP 的差异

MLP 的基本单元通常可以写为：`y = W σ(x) + b`。其中 `σ` 是 ReLU、GELU、SiLU 等固定激活函数，网络学习的主要是线性权重 `W`。KAN 的思路不同，它把每条边上的映射定义为可学习的一维函数，节点主要负责汇总入边输出。简化表示为：

`f(x) = Σ_e φ_e(x_e),  φ_e(x)=Σ_i c_i B_i(x)`

其中 `B_i(x)` 是 B-spline 基函数，`c_i` 是可学习系数。本文复现的 `KANLinear` 同时保留基础激活分支和样条分支：基础分支用于稳定训练，样条分支用于学习更灵活的非线性边函数。

![图1 MLP 与 KAN 的非线性建模差异]({rel(figs["kan_vs_mlp"])})

图 1 来源说明：作者根据原始手绘草图 `{HANDDRAWN_PDF.name}`、KAN 论文思想和本文复现代码重新绘制。

### 2.2 U-KAN 分割网络

本文复现的 U-KAN 使用三层卷积编码器提取局部纹理特征，然后将中高层特征通过 PatchEmbed 转换为 token 序列，并输入 KANBlock。KANBlock 内部包含 LayerNorm、KANLayer、深度可分离卷积、BatchNorm 与 ReLU，并通过残差连接保持训练稳定。解码器部分逐级上采样，并将对应尺度的编码器特征通过 skip connection 加回。

![图2 本文复现的 U-KAN 医学图像分割结构]({rel(figs["ukan_architecture"])})

图 2 来源说明：作者根据 U-KAN 官方源码结构与本文 `src/ukan_course/models/ukan.py` 复现代码整理绘制。

在消融模型 U-KAN(no-KAN) 中，本文保持 U-KAN 的整体结构、输入尺度、训练参数和解码器完全一致，仅将 KANLinear 替换为普通 `nn.Linear`。因此，U-KAN 与 no-KAN 的性能差异可以更直接地反映 KAN 模块本身的贡献，而不是网络深度、跳跃连接或数据处理差异。

### 2.3 Attention-U-KAN 改进尝试

U-KAN 的 skip connection 直接将编码器细节与解码器语义特征相加。对于医学图像，低层特征中既包含边界纹理，也可能包含噪声和背景结构。因此本文尝试在 skip fusion 后加入轻量通道-空间注意力模块。该模块先通过全局平均池化和 1×1 卷积生成通道权重，再通过通道维平均图、最大图和 7×7 卷积生成空间权重，以此重标定融合特征。

![图3 Attention-U-KAN 的 skip fusion 增强模块]({rel(figs["attention_module"])})

图 3 来源说明：作者根据本文 `attention.py` 与 `attention_ukan.py` 代码自绘。

需要说明的是，本文没有预设 Attention-U-KAN 一定优于 U-KAN，而是把它作为一个合理改进方向进行验证。如果实验结果不提升，仍然可以作为有效的消融结论：简单注意力模块是否适合 U-KAN，需要由数据和指标决定。

## 3 数据集、实验平台与实验设计

### 3.1 数据集与预处理

本文使用两个公开医学图像分割数据集。BUSI 数据集包含乳腺超声图像及病灶 mask，原始数据中部分病例存在多个 mask，本文在预处理阶段将同一图像对应的多个 mask 取并集合并，统一为单通道二值 mask。CVC-ClinicDB 数据集用于结肠镜息肉分割，图像和 mask 一一对应。两个数据集均被整理为 `data/processed/<dataset>/images` 与 `data/processed/<dataset>/masks` 格式，并通过固定 split 文件保存训练集和验证集划分。

{markdown_table(["数据集", "图像数", "宽度范围", "高度范围", "平均 mask 占比", "处理说明"], tables["dataset"])}

表 1 来源说明：由本文 `dataset_report.py` 对整理后的本地数据统计生成。

![图4 BUSI 样本可视化](../experiments/figures/busi_samples.png)

图 4 来源说明：由本文脚本从整理后的 BUSI 数据中抽样生成，原始数据来自公开 BUSI 数据集。

![图5 CVC-ClinicDB 样本可视化](../experiments/figures/cvc_samples.png)

图 5 来源说明：由本文脚本从整理后的 CVC 数据中抽样生成，原始数据来自公开 CVC-ClinicDB 数据集。

### 3.2 实验平台与训练设置

本文实验在 Windows + Anaconda PyTorch 环境完成，训练和评估均使用同一块 NVIDIA RTX 4080 Laptop GPU。所有主实验使用 256×256 输入分辨率、batch size 8、100 epoch、Adam 优化器和 CosineAnnealingLR 调度器。损失函数为 BCEWithLogits 与 DiceLoss 的组合：

`Loss = 0.5 × BCEWithLogits(logits, mask) + (1 - Dice_soft)`

最终评估指标统一使用像素级混淆矩阵计算：

`IoU = TP / (TP + FP + FN)`  
`Dice = 2TP / (2TP + FP + FN)`  
`Precision = TP / (TP + FP)`  
`Recall = TP / (TP + FN)`  
`Specificity = TN / (TN + FP)`

{markdown_table(["项目", "设置", "说明"], tables["platform"])}

表 2 来源说明：由 `doctor_env.py`、YAML 配置文件和训练脚本记录整理。

### 3.3 对比实验设计

本文主实验包含 BUSI 与 CVC 两个数据集，每个数据集运行四组模型：U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN。U-Net 是经典 CNN 分割基线；no-KAN 是与 U-KAN 结构最接近的消融模型；U-KAN 是本文主复现模型；Attention-U-KAN 是本文改进尝试。实验流程如图 6 所示。

![图6 实验流程与课程论文产出链路]({rel(figs["workflow"])})

图 6 来源说明：作者根据本文数据处理、训练、评估和论文产出流程自绘。

为了回应训练过程中发现的问题，本文还设计了两个补充实验。其一，针对 BUSI no-KAN 曲线中局部跳升现象，补充 seed 6142 下的 U-KAN 与 no-KAN 对照，以判断 KAN 收益是否只是单一划分偶然结果。其二，针对 CVC 若干模型 best epoch 靠后的现象，从 CVC U-KAN 的 best checkpoint 出发继续训练 50 epoch，判断主复现模型是否仍未收敛。

## 4 实验结果与分析

### 4.1 BUSI 主实验结果

{markdown_table(result_headers, tables["busi_results"])}

表 3 来源说明：由 `evaluate.py` 在 BUSI 验证集上统一评估生成；推理时间为单图平均耗时。

BUSI 结果显示，U-KAN 取得最佳 IoU 与 Dice，分别为 {stats["busi_ukan_iou"]} 和 {stats["busi_ukan_dice"]}。与 U-KAN(no-KAN) 相比，U-KAN 的 IoU 提升 {stats["busi_gain_nokan_iou"]}，说明在相同 U 型框架中引入 KANLinear 能够提升病灶区域重叠质量。与 U-Net 相比，U-KAN 的 IoU 提升 {stats["busi_gain_unet_iou"]}，尤其 Recall 从 U-Net 的较低水平提升到 U-KAN 的 0.8118，说明 U-KAN 对病灶区域的召回能力更强。

Attention-U-KAN 在 BUSI 上低于原始 U-KAN，IoU 差距为 {stats["busi_att_gap"]}。这说明简单地在 skip fusion 后加入通道-空间注意力并不必然有效。BUSI 图像具有明显噪声和弱边界，注意力模块可能在突出高响应区域的同时抑制了部分模糊病灶边缘。

![图7 BUSI 训练曲线](../experiments/figures/busi_training_curves.png)

图 7 来源说明：由 `plot_training_curves.py` 根据训练日志生成。

![图8 BUSI 预测结果对比](../experiments/figures/busi_prediction_comparison.png)

图 8 来源说明：由 `visualize_predictions.py` 根据验证集预测结果生成，包含原图、真值、预测和误差图。

### 4.2 CVC 主实验结果

{markdown_table(result_headers, tables["cvc_results"])}

表 4 来源说明：由 `evaluate.py` 在 CVC 验证集上统一评估生成；推理时间为单图平均耗时。

CVC 结果中，U-KAN 同样取得主实验最高 IoU 和 Dice，分别为 {stats["cvc_ukan_iou"]} 和 {stats["cvc_ukan_dice"]}。相对 no-KAN，U-KAN 的 IoU 提升 {stats["cvc_gain_nokan_iou"]}；相对 U-Net，IoU 提升 {stats["cvc_gain_unet_iou"]}。相比 BUSI，CVC 中 U-KAN 对 U-Net 的提升幅度较小，可能是因为 CVC 图像分辨率固定、目标结构更规则，传统 U-Net 已经可以获得较强基线。

Attention-U-KAN 在 CVC 上低于 U-KAN，IoU 差距为 {stats["cvc_att_gap"]}。这一结果与 BUSI 一致，说明本文采用的轻量注意力插入方式没有形成稳定收益。值得注意的是，Attention-U-KAN 的 Precision 较高，但 Recall 偏低，说明它更倾向于保守预测，减少误分背景的同时可能漏掉部分真实息肉区域。

![图9 CVC 训练曲线](../experiments/figures/cvc_training_curves.png)

图 9 来源说明：由 `plot_training_curves.py` 根据训练日志生成。

![图10 CVC 预测结果对比](../experiments/figures/cvc_prediction_comparison.png)

图 10 来源说明：由 `visualize_predictions.py` 根据验证集预测结果生成，包含原图、真值、预测和误差图。

### 4.3 训练诊断与补充实验

训练曲线显示，BUSI 的 no-KAN 在部分 epoch 出现验证 IoU 局部跳升，而 CVC 中部分模型的 best epoch 靠近训练末端。为避免只凭单次曲线下结论，本文使用训练诊断表和补充实验进行复核。

{markdown_table(["实验", "best epoch", "日志 best IoU", "最终 epoch IoU", "后 10 epoch 变化", "evaluate IoU"], tables["diagnostics"])}

表 5 来源说明：由 `analyze_training_logs.py` 汇总训练日志和最终评估结果生成。前 8 个主实验的旧训练日志曾使用 batch 平均指标，最终比较以 `evaluate.py` 全验证集混淆矩阵结果为准。

BUSI seed 6142 补充实验如下：

{markdown_table(["模型", "seed2981 IoU", "seed2981 Dice", "seed6142 IoU", "seed6142 Dice", "两 seed 平均 IoU", "两 seed 平均 Dice"], tables["seed_stability"])}

表 6 来源说明：由 BUSI seed2981 与 seed6142 的 `evaluate.py` 结果整理生成。

seed 6142 中，U-KAN 仍高于 no-KAN；两 seed 平均后，U-KAN 的 IoU 为 0.6515，高于 no-KAN 的 0.6214。这说明 KAN 模块收益不是 seed 2981 的单次偶然现象。同时，seed 6142 的绝对指标整体低于 seed 2981，表明 BUSI 对随机划分和初始化较敏感，论文结论应强调“复现与课程资源条件下的验证”，而不是宣称完全复刻官方多 seed 最优结果。

CVC 继续训练补充实验如下：

{markdown_table(["实验", "设置", "epoch", "best epoch", "IoU", "Dice", "后 10 epoch IoU 变化"], tables["cvc_continue"])}

表 7 来源说明：由 CVC U-KAN 原始训练与继续训练的日志、评估结果整理生成。

从原始 U-KAN best checkpoint 继续训练 50 epoch 后，IoU 为 0.7847，低于主实验的 0.7874，差距为 {stats["cvc_ft_gap"]}。因此，当前证据不支持“CVC 结果偏低只是因为没有继续训练”的解释。更合理的原因包括：本文主实验为 100 epoch 课程复现实验，官方通常使用更长训练轮数、多 seed 平均、不同 checkpoint 或不同数据处理细节。

### 4.4 过程记录与工程复现

为满足课程对过程记录和源码提交的要求，本文保留了训练、评估、预测 mask 生成和结果汇总截图，并在 GitHub 仓库中提供可复现的脚本、配置文件和 README。

![图11 训练、评估与结果生成过程截图]({rel(figs["screenshots"])})

图 11 来源说明：作者本地运行训练、评估和可视化脚本时截取，本文整理为拼图。

## 5 讨论

第一，U-KAN 的优势主要体现在 KAN 模块本身，而不是简单的参数量堆叠。BUSI 中 no-KAN 的参数量约 2.76M，U-KAN 约 6.36M，U-Net 约 7.76M。虽然 U-KAN 参数量低于 U-Net，但仍取得更高 IoU 和 Dice，说明其 tokenized KANBlock 对医学图像分割具有有效表达能力。CVC 中 U-KAN 对 U-Net 的提升较小，但仍保持最高主指标，说明其收益与数据集难度和目标形态有关。

第二，KAN 带来了明显推理代价。U-Net 单图推理时间约 3.6–3.7 ms，而 U-KAN 系列约 37–46 ms。这与 KANLinear 的 B-spline 基函数计算、token 展平和高维线性变换有关。因此在实际医学部署中，需要在分割精度和推理效率之间权衡。对于实时内镜场景，U-KAN 可能需要剪枝、蒸馏或轻量化；对于离线超声辅助分析，较高推理时间更容易接受。

第三，Attention-U-KAN 的负结果是有价值的。很多课程论文容易把“加入注意力机制”直接写成提升点，但本文实验显示简单插入注意力并不稳定。可能原因包括：数据规模有限导致注意力模块过拟合；通道-空间注意力过度抑制模糊边界；插入位置不够精细；U-KAN 本身已经在中高层通过 KANBlock 提升了非线性表达，简单注意力的边际收益有限。后续若继续改进，可以尝试只在高层语义 skip 上加入注意力，或结合边界损失、多尺度监督和不确定性建模。

第四，本文对训练日志口径进行了修正。早期主实验的训练曲线用于观察收敛趋势，但旧日志中的验证 IoU/Dice 采用 batch 平均；最终论文表格全部采用 `evaluate.py` 的全验证集像素级混淆矩阵口径。补充实验之后，`train.py` 已调整为与 `evaluate.py` 一致的全局统计方式。因此，新旧训练曲线不应直接比较绝对值，最终模型比较应以统一评估表为准。

第五，本文仍存在局限。受课程时间和算力约束，CVC 没有完成官方推荐的 400 epoch 和三 seed 完整平均；BUSI 只补充了一个额外 seed；当前评估集来自固定验证划分，不是完全独立外部测试集。尽管如此，本文完成了从数据整理、模型复现、消融对比、补充实验、可视化到工程文档的完整流程，足以支撑课程论文中的算法复现和独立分析。

## 6 课程要求与加分项完成情况

{markdown_table(["加分/要求项", "完成方式", "论文或项目位置"], [
["GitHub 源码链接", f"完整工程已配置远程仓库 {GITHUB_URL}", "README 与本文来源说明"],
["详细 README", "包含环境、数据准备、训练、评估、绘图和结果汇总命令", "README.md"],
["项目结构", "configs/scripts/src/tests/experiments/paper 分层组织", "GitHub 仓库"],
["数据整理", "BUSI/CVC 公开数据统一清洗、mask 合并、固定 split", "scripts/prepare_*.py 与 data/splits"],
["自绘图", "KAN 原理图、U-KAN 结构图、Attention 模块图、实验流程图", "paper/figures"],
["图表编号与来源", "所有图表均编号并标注生成脚本或数据来源", "论文正文"],
["过程记录", "训练与评估截图、训练日志、metrics.csv、预测 mask", "课程论文要求/图片 与 experiments/results"],
["独立思考", "针对曲线异常和 CVC 收敛性追加 seed 与继续训练实验", "表 5–7 与讨论章节"],
] )}

## 7 结论

本文完成了 U-KAN 医学图像分割算法的课程复现与扩展实验。通过 BUSI 和 CVC 两个数据集的主实验，U-KAN 均取得当前实验矩阵中的最佳 IoU 和 Dice；与 no-KAN 的对比说明 KANLinear 对分割性能具有实际贡献；与 U-Net 的对比说明 U-KAN 在精度上具有优势，但推理速度明显较慢。Attention-U-KAN 的实验结果未超过原始 U-KAN，提示简单注意力机制并非稳定改进方向。

补充实验进一步增强了结论可信度。BUSI seed 6142 显示 U-KAN 相对 no-KAN 的优势仍然存在，但绝对指标受随机划分影响较大；CVC 继续训练实验没有超过原始 100 epoch best checkpoint，说明当前 U-KAN 结果与官方参考差距不能简单归因于训练不够。总体而言，本文不仅完成了算法复现，还通过消融实验、补充实验和负结果讨论体现了独立分析能力。最终项目形成了可运行代码、清晰 README、数据处理脚本、实验日志、图表资产和规范课程论文，满足课程方向（一）及多个加分项要求。

## 参考文献

[1] Li C., Liu X., Li W., et al. U-KAN Makes Strong Backbone for Medical Image Segmentation and Generation. arXiv preprint, 2024.  
[2] Liu Z., Wang Y., Vaidya S., et al. KAN: Kolmogorov-Arnold Networks. arXiv:2404.19756, 2024.  
[3] Ronneberger O., Fischer P., Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation. MICCAI, 2015.  
[4] Woo S., Park J., Lee J.-Y., Kweon I. S. CBAM: Convolutional Block Attention Module. ECCV, 2018.  
[5] Al-Dhabyani W., Gomaa M., Khaled H., Fahmy A. Dataset of Breast Ultrasound Images. Data in Brief, 2020.  
[6] Bernal J., Sánchez F. J., Fernández-Esparrach G., et al. WM-DOVA maps for accurate polyp highlighting in colonoscopy. Computerized Medical Imaging and Graphics, 2015.  
[7] CUHK-AIM-Group. U-KAN official GitHub repository. https://github.com/CUHK-AIM-Group/U-KAN  
[8] Liu Z. pykan official GitHub repository. https://github.com/KindXiaoming/pykan  
[9] A Review of KANs for Medical Image Analysis. 本地课程材料中的 KAN 医学图像分析综述资料。
"""
    MD_OUT.write_text(md, encoding="utf-8")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size: float = 12, bold: bool = False, color: str | None = None, east_asia: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc: Document, text: str, first_line: bool = True, align: int | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, 12)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 14, True, east_asia="黑体")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, 10.5)


def add_figure(doc: Document, path: Path, caption: str, source: str, width_cm: float = 15.5) -> None:
    if not path.exists():
        add_para(doc, f"（缺少图片文件：{path}）", first_line=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, f"{caption}\n来源：{source}")


def add_table(doc: Document, title: str, headers: list[str], rows: list[list[str]], source: str, font_size: float = 9.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, 10.5, True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        header_cells[i].text = h
        set_cell_shading(header_cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, font_size)
    add_caption(doc, f"来源：{source}")


def build_docx(figs: dict[str, Path], tables: dict[str, list[list[str]]], stats: dict[str, str]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run(TITLE)
    set_run_font(r, 22, True, east_asia="黑体")
    for item in [
        "课程方向：方向（一）复现任意人工智能算法",
        "课程：人工智能",
        "作者：蔡雪峰",
        "日期：2026 年 6 月",
        f"项目 GitHub：{GITHUB_URL}",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(item)
        set_run_font(run, 12)
    doc.add_page_break()

    add_heading(doc, "来源说明", 1)
    add_para(
        doc,
        "本文围绕 U-KAN 在二维医学图像分割任务中的复现与改进展开。算法来源主要包括 U-KAN 论文与官方开源项目、Kolmogorov-Arnold Network（KAN）论文及 pykan 项目、U-Net 经典医学图像分割框架、通道-空间注意力机制相关论文，以及 BUSI 与 CVC-ClinicDB 两个公开医学图像分割数据集。项目代码为课程论文重新组织的 PyTorch 工程，未直接复制官方训练脚本；官方源码主要用于核对模型结构、训练超参数和结果解释。",
    )
    add_para(
        doc,
        f"本项目使用的全部实验指标来自本地训练和 evaluate.py 脚本评估，未由大语言模型生成或改写。本文写作、结构组织、语言润色和图表说明使用 OpenAI Codex 辅助完成；实验方案、训练数据、运行截图、结果表格和结论均基于本地工程记录。原始手绘草图文件为 {HANDDRAWN_PDF.name}，本文进一步绘制了规范化 KAN 原理图、U-KAN 结构图、Attention-U-KAN 模块图和实验流程图，并在每幅图下标明来源。",
    )

    add_heading(doc, "摘要", 1)
    add_para(
        doc,
        f"医学图像分割需要在有限标注数据条件下准确定位病灶或器官边界。传统 U-Net 结构具有清晰的编码器-解码器和跳跃连接，但中高层非线性表达通常依赖卷积或普通 MLP。U-KAN 将 Kolmogorov-Arnold Network 中“边上可学习一维函数”的思想引入 U 型分割网络，在 token 化特征上使用 KANBlock 建模复杂非线性关系。本文选择课程方向（一）“复现任意人工智能算法”，在本地 PyTorch 环境中复现 U-KAN 的医学图像分割流程，并在此基础上尝试加入轻量通道-空间注意力模块，构建 Attention-U-KAN。实验使用 BUSI 乳腺超声数据集和 CVC-ClinicDB 结肠镜息肉数据集，统一进行数据整理、固定随机划分、训练、评估和可视化。结果显示，U-KAN 在 BUSI 上取得 IoU={stats['busi_ukan_iou']}、Dice={stats['busi_ukan_dice']}，在 CVC 上取得 IoU={stats['cvc_ukan_iou']}、Dice={stats['cvc_ukan_dice']}，均为主实验最佳。补充实验表明，BUSI 中 KAN 模块收益不是单一 seed 偶然现象；CVC 继续训练实验显示当前 U-KAN 主模型基本达到平台期。Attention-U-KAN 未稳定超过原始 U-KAN，说明简单注意力插入并不必然提升 KAN 分割网络。",
    )
    add_para(doc, "关键词：医学图像分割；U-KAN；Kolmogorov-Arnold Network；U-Net；注意力机制；算法复现", first_line=False)

    add_heading(doc, "1 引言", 1)
    for text in [
        "医学图像分割是计算机辅助诊断中的基础任务之一。对于乳腺超声、结肠镜等图像，准确分割病灶区域可以辅助医生估计病灶大小、形态和边界，为后续诊断和随访提供量化依据。与自然图像相比，医学图像常具有样本数量有限、标注成本高、边界模糊、目标面积小和图像噪声强等特点，因此模型不仅需要较强的局部纹理提取能力，也需要在中高层语义上表达复杂非线性关系。",
        "U-Net 是医学图像分割中最经典的卷积网络结构。它通过编码器提取多尺度语义特征，通过解码器逐步恢复空间分辨率，并利用跳跃连接弥补上采样过程中的细节损失。此类结构简单、稳定、易训练，但在复杂组织边界和弱对比目标上仍可能受到卷积感受野和普通非线性变换的限制。KAN 提供了另一种建模思路：传统 MLP 通常在节点上放置固定激活函数，而 KAN 将可学习的一维函数放在网络边上，并用样条基函数进行参数化。",
        "U-KAN 的核心动机是把 KAN 的非线性表达能力嵌入 U 型医学图像分割架构。它保留 U-Net 的局部卷积编码和跳跃连接，同时在更高层的 token 化特征上使用 KANBlock 替代普通 MLP，以增强全局语义和非线性映射能力。本文选择 U-KAN 作为复现对象，是因为它既符合“复现一个人工智能算法”的课程要求，又能自然扩展出医学图像分割数据整理、模型消融、对比实验、可视化和工程化复现等加分项。",
        "本文的主要工作包括：第一，重新组织 U-KAN 分割工程，完成 BUSI 与 CVC 数据集预处理、固定划分、训练和评估；第二，实现 U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN 四类模型，形成主实验矩阵；第三，围绕训练曲线异常和 CVC 收敛性疑问设计补充实验；第四，生成完整图表、运行截图、README 和论文正文，从实验事实出发讨论复现效果与局限。",
    ]:
        add_para(doc, text)

    add_heading(doc, "2 算法原理与模型构建", 1)
    add_heading(doc, "2.1 KAN 与 MLP 的差异", 2)
    add_para(doc, "MLP 的基本单元通常可以写为 y = Wσ(x) + b，其中 σ 是 ReLU、GELU、SiLU 等固定激活函数，网络学习的主要是线性权重 W。KAN 的思路不同，它把每条边上的映射定义为可学习的一维函数，节点主要负责汇总入边输出。简化表示为：φe(x)=ΣciBi(x)，其中 Bi(x) 是 B-spline 基函数，ci 是可学习系数。本文复现的 KANLinear 同时保留基础激活分支和样条分支：基础分支用于稳定训练，样条分支用于学习更灵活的非线性边函数。")
    add_figure(doc, figs["kan_vs_mlp"], "图 1 MLP 与 KAN 的非线性建模差异", f"作者根据原始手绘草图 {HANDDRAWN_PDF.name}、KAN 论文思想和本文复现代码重新绘制。", 16.0)

    add_heading(doc, "2.2 U-KAN 分割网络", 2)
    add_para(doc, "本文复现的 U-KAN 使用三层卷积编码器提取局部纹理特征，然后将中高层特征通过 PatchEmbed 转换为 token 序列，并输入 KANBlock。KANBlock 内部包含 LayerNorm、KANLayer、深度可分离卷积、BatchNorm 与 ReLU，并通过残差连接保持训练稳定。解码器部分逐级上采样，并将对应尺度的编码器特征通过 skip connection 加回。")
    add_figure(doc, figs["ukan_architecture"], "图 2 本文复现的 U-KAN 医学图像分割结构", "作者根据 U-KAN 官方源码结构与本文 src/ukan_course/models/ukan.py 复现代码整理绘制。", 16.0)
    add_para(doc, "在消融模型 U-KAN(no-KAN) 中，本文保持 U-KAN 的整体结构、输入尺度、训练参数和解码器完全一致，仅将 KANLinear 替换为普通 nn.Linear。因此，U-KAN 与 no-KAN 的性能差异可以更直接地反映 KAN 模块本身的贡献，而不是网络深度、跳跃连接或数据处理差异。")

    add_heading(doc, "2.3 Attention-U-KAN 改进尝试", 2)
    add_para(doc, "U-KAN 的 skip connection 直接将编码器细节与解码器语义特征相加。对于医学图像，低层特征中既包含边界纹理，也可能包含噪声和背景结构。因此本文尝试在 skip fusion 后加入轻量通道-空间注意力模块。该模块先通过全局平均池化和 1×1 卷积生成通道权重，再通过通道维平均图、最大图和 7×7 卷积生成空间权重，以此重标定融合特征。")
    add_figure(doc, figs["attention_module"], "图 3 Attention-U-KAN 的 skip fusion 增强模块", "作者根据本文 attention.py 与 attention_ukan.py 代码自绘。", 16.0)
    add_para(doc, "需要说明的是，本文没有预设 Attention-U-KAN 一定优于 U-KAN，而是把它作为一个合理改进方向进行验证。如果实验结果不提升，仍然可以作为有效的消融结论：简单注意力模块是否适合 U-KAN，需要由数据和指标决定。")

    add_heading(doc, "3 数据集、实验平台与实验设计", 1)
    add_heading(doc, "3.1 数据集与预处理", 2)
    add_para(doc, "本文使用两个公开医学图像分割数据集。BUSI 数据集包含乳腺超声图像及病灶 mask，原始数据中部分病例存在多个 mask，本文在预处理阶段将同一图像对应的多个 mask 取并集合并，统一为单通道二值 mask。CVC-ClinicDB 数据集用于结肠镜息肉分割，图像和 mask 一一对应。两个数据集均被整理为 data/processed/<dataset>/images 与 data/processed/<dataset>/masks 格式，并通过固定 split 文件保存训练集和验证集划分。")
    add_table(doc, "表 1 数据集统计与预处理说明", ["数据集", "图像数", "宽度范围", "高度范围", "平均 mask 占比", "处理说明"], tables["dataset"], "由本文 dataset_report.py 对整理后的本地数据统计生成。", 8.5)
    add_figure(doc, figs["busi_samples"], "图 4 BUSI 样本可视化", "由本文脚本从整理后的 BUSI 数据中抽样生成，原始数据来自公开 BUSI 数据集。", 15.8)
    add_figure(doc, figs["cvc_samples"], "图 5 CVC-ClinicDB 样本可视化", "由本文脚本从整理后的 CVC 数据中抽样生成，原始数据来自公开 CVC-ClinicDB 数据集。", 15.8)

    add_heading(doc, "3.2 实验平台与训练设置", 2)
    add_para(doc, "本文实验在 Windows + Anaconda PyTorch 环境完成，训练和评估均使用同一块 NVIDIA RTX 4080 Laptop GPU。所有主实验使用 256×256 输入分辨率、batch size 8、100 epoch、Adam 优化器和 CosineAnnealingLR 调度器。损失函数为：Loss = 0.5×BCEWithLogits + (1 - Dice_soft)。最终评估指标统一使用像素级混淆矩阵计算 IoU、Dice、Precision、Recall 和 Specificity。")
    add_table(doc, "表 2 实验平台与统一训练设置", ["项目", "设置", "说明"], tables["platform"], "由 doctor_env.py、YAML 配置文件和训练脚本记录整理。", 9.5)

    add_heading(doc, "3.3 对比实验设计", 2)
    add_para(doc, "本文主实验包含 BUSI 与 CVC 两个数据集，每个数据集运行四组模型：U-Net、U-KAN(no-KAN)、U-KAN 和 Attention-U-KAN。U-Net 是经典 CNN 分割基线；no-KAN 是与 U-KAN 结构最接近的消融模型；U-KAN 是本文主复现模型；Attention-U-KAN 是本文改进尝试。")
    add_figure(doc, figs["workflow"], "图 6 实验流程与课程论文产出链路", "作者根据本文数据处理、训练、评估和论文产出流程自绘。", 16.0)
    add_para(doc, "为了回应训练过程中发现的问题，本文还设计了两个补充实验。其一，针对 BUSI no-KAN 曲线中局部跳升现象，补充 seed 6142 下的 U-KAN 与 no-KAN 对照，以判断 KAN 收益是否只是单一划分偶然结果。其二，针对 CVC 若干模型 best epoch 靠后的现象，从 CVC U-KAN 的 best checkpoint 出发继续训练 50 epoch，判断主复现模型是否仍未收敛。")

    result_headers = ["模型", "IoU", "Dice", "Precision", "Recall", "Specificity", "参数量", "推理 ms/图"]
    add_heading(doc, "4 实验结果与分析", 1)
    add_heading(doc, "4.1 BUSI 主实验结果", 2)
    add_table(doc, "表 3 BUSI 主实验结果", result_headers, tables["busi_results"], "由 evaluate.py 在 BUSI 验证集上统一评估生成；推理时间为单图平均耗时。", 8.2)
    add_para(doc, f"BUSI 结果显示，U-KAN 取得最佳 IoU 与 Dice，分别为 {stats['busi_ukan_iou']} 和 {stats['busi_ukan_dice']}。与 U-KAN(no-KAN) 相比，U-KAN 的 IoU 提升 {stats['busi_gain_nokan_iou']}，说明在相同 U 型框架中引入 KANLinear 能够提升病灶区域重叠质量。与 U-Net 相比，U-KAN 的 IoU 提升 {stats['busi_gain_unet_iou']}，尤其 Recall 从 U-Net 的较低水平提升到 U-KAN 的 0.8118，说明 U-KAN 对病灶区域的召回能力更强。Attention-U-KAN 在 BUSI 上低于原始 U-KAN，说明简单地在 skip fusion 后加入通道-空间注意力并不必然有效。")
    add_figure(doc, figs["busi_curves"], "图 7 BUSI 训练曲线", "由 plot_training_curves.py 根据训练日志生成。", 16.0)
    add_figure(doc, figs["busi_preds"], "图 8 BUSI 预测结果对比", "由 visualize_predictions.py 根据验证集预测结果生成，包含原图、真值、预测和误差图。", 16.0)

    add_heading(doc, "4.2 CVC 主实验结果", 2)
    add_table(doc, "表 4 CVC 主实验结果", result_headers, tables["cvc_results"], "由 evaluate.py 在 CVC 验证集上统一评估生成；推理时间为单图平均耗时。", 8.2)
    add_para(doc, f"CVC 结果中，U-KAN 同样取得主实验最高 IoU 和 Dice，分别为 {stats['cvc_ukan_iou']} 和 {stats['cvc_ukan_dice']}。相对 no-KAN，U-KAN 的 IoU 提升 {stats['cvc_gain_nokan_iou']}；相对 U-Net，IoU 提升 {stats['cvc_gain_unet_iou']}。相比 BUSI，CVC 中 U-KAN 对 U-Net 的提升幅度较小，可能是因为 CVC 图像分辨率固定、目标结构更规则，传统 U-Net 已经可以获得较强基线。Attention-U-KAN 在 CVC 上低于 U-KAN，其 Precision 较高但 Recall 偏低，说明模型更倾向于保守预测。")
    add_figure(doc, figs["cvc_curves"], "图 9 CVC 训练曲线", "由 plot_training_curves.py 根据训练日志生成。", 16.0)
    add_figure(doc, figs["cvc_preds"], "图 10 CVC 预测结果对比", "由 visualize_predictions.py 根据验证集预测结果生成，包含原图、真值、预测和误差图。", 16.0)

    add_heading(doc, "4.3 训练诊断与补充实验", 2)
    add_para(doc, "训练曲线显示，BUSI 的 no-KAN 在部分 epoch 出现验证 IoU 局部跳升，而 CVC 中部分模型的 best epoch 靠近训练末端。为避免只凭单次曲线下结论，本文使用训练诊断表和补充实验进行复核。")
    add_table(doc, "表 5 训练日志诊断", ["实验", "best epoch", "日志 best IoU", "最终 epoch IoU", "后 10 epoch 变化", "evaluate IoU"], tables["diagnostics"], "由 analyze_training_logs.py 汇总训练日志和最终评估结果生成。前 8 个主实验的旧训练日志曾使用 batch 平均指标，最终比较以 evaluate.py 全验证集混淆矩阵结果为准。", 8.2)
    add_table(doc, "表 6 BUSI seed 稳定性补充实验", ["模型", "seed2981 IoU", "seed2981 Dice", "seed6142 IoU", "seed6142 Dice", "两 seed 平均 IoU", "两 seed 平均 Dice"], tables["seed_stability"], "由 BUSI seed2981 与 seed6142 的 evaluate.py 结果整理生成。", 8.0)
    add_para(doc, "seed 6142 中，U-KAN 仍高于 no-KAN；两 seed 平均后，U-KAN 的 IoU 为 0.6515，高于 no-KAN 的 0.6214。这说明 KAN 模块收益不是 seed 2981 的单次偶然现象。同时，seed 6142 的绝对指标整体低于 seed 2981，表明 BUSI 对随机划分和初始化较敏感，论文结论应强调“复现与课程资源条件下的验证”。")
    add_table(doc, "表 7 CVC U-KAN 继续训练补充实验", ["实验", "设置", "epoch", "best epoch", "IoU", "Dice", "后 10 epoch IoU 变化"], tables["cvc_continue"], "由 CVC U-KAN 原始训练与继续训练的日志、评估结果整理生成。", 8.0)
    add_para(doc, f"从原始 U-KAN best checkpoint 继续训练 50 epoch 后，IoU 为 0.7847，低于主实验的 0.7874，差距为 {stats['cvc_ft_gap']}。因此，当前证据不支持“CVC 结果偏低只是因为没有继续训练”的解释。更合理的原因包括：本文主实验为 100 epoch 课程复现实验，官方通常使用更长训练轮数、多 seed 平均、不同 checkpoint 或不同数据处理细节。")

    add_heading(doc, "4.4 过程记录与工程复现", 2)
    add_para(doc, "为满足课程对过程记录和源码提交的要求，本文保留了训练、评估、预测 mask 生成和结果汇总截图，并在 GitHub 仓库中提供可复现的脚本、配置文件和 README。")
    add_figure(doc, figs["screenshots"], "图 11 训练、评估与结果生成过程截图", "作者本地运行训练、评估和可视化脚本时截取，本文整理为拼图。", 16.0)

    add_heading(doc, "5 讨论", 1)
    for text in [
        "第一，U-KAN 的优势主要体现在 KAN 模块本身，而不是简单的参数量堆叠。BUSI 中 no-KAN 的参数量约 2.76M，U-KAN 约 6.36M，U-Net 约 7.76M。虽然 U-KAN 参数量低于 U-Net，但仍取得更高 IoU 和 Dice，说明其 tokenized KANBlock 对医学图像分割具有有效表达能力。CVC 中 U-KAN 对 U-Net 的提升较小，但仍保持最高主指标，说明其收益与数据集难度和目标形态有关。",
        "第二，KAN 带来了明显推理代价。U-Net 单图推理时间约 3.6–3.7 ms，而 U-KAN 系列约 37–46 ms。这与 KANLinear 的 B-spline 基函数计算、token 展平和高维线性变换有关。因此在实际医学部署中，需要在分割精度和推理效率之间权衡。对于实时内镜场景，U-KAN 可能需要剪枝、蒸馏或轻量化；对于离线超声辅助分析，较高推理时间更容易接受。",
        "第三，Attention-U-KAN 的负结果是有价值的。很多课程论文容易把“加入注意力机制”直接写成提升点，但本文实验显示简单插入注意力并不稳定。可能原因包括：数据规模有限导致注意力模块过拟合；通道-空间注意力过度抑制模糊边界；插入位置不够精细；U-KAN 本身已经在中高层通过 KANBlock 提升了非线性表达，简单注意力的边际收益有限。",
        "第四，本文对训练日志口径进行了修正。早期主实验的训练曲线用于观察收敛趋势，但旧日志中的验证 IoU/Dice 采用 batch 平均；最终论文表格全部采用 evaluate.py 的全验证集像素级混淆矩阵口径。补充实验之后，train.py 已调整为与 evaluate.py 一致的全局统计方式。因此，新旧训练曲线不应直接比较绝对值，最终模型比较应以统一评估表为准。",
        "第五，本文仍存在局限。受课程时间和算力约束，CVC 没有完成官方推荐的 400 epoch 和三 seed 完整平均；BUSI 只补充了一个额外 seed；当前评估集来自固定验证划分，不是完全独立外部测试集。尽管如此，本文完成了从数据整理、模型复现、消融对比、补充实验、可视化到工程文档的完整流程，足以支撑课程论文中的算法复现和独立分析。",
    ]:
        add_para(doc, text)

    add_heading(doc, "6 课程要求与加分项完成情况", 1)
    bonus_rows = [
        ["GitHub 源码链接", f"完整工程已配置远程仓库 {GITHUB_URL}", "README 与本文来源说明"],
        ["详细 README", "包含环境、数据准备、训练、评估、绘图和结果汇总命令", "README.md"],
        ["项目结构", "configs/scripts/src/tests/experiments/paper 分层组织", "GitHub 仓库"],
        ["数据整理", "BUSI/CVC 公开数据统一清洗、mask 合并、固定 split", "scripts/prepare_*.py 与 data/splits"],
        ["自绘图", "KAN 原理图、U-KAN 结构图、Attention 模块图、实验流程图", "paper/figures"],
        ["图表编号与来源", "所有图表均编号并标注生成脚本或数据来源", "论文正文"],
        ["过程记录", "训练与评估截图、训练日志、metrics.csv、预测 mask", "课程论文要求/图片 与 experiments/results"],
        ["独立思考", "针对曲线异常和 CVC 收敛性追加 seed 与继续训练实验", "表 5–7 与讨论章节"],
    ]
    add_table(doc, "表 8 课程要求与加分项落实情况", ["加分/要求项", "完成方式", "论文或项目位置"], bonus_rows, "根据课程论文要求与本文工程材料整理。", 8.0)

    add_heading(doc, "7 结论", 1)
    for text in [
        "本文完成了 U-KAN 医学图像分割算法的课程复现与扩展实验。通过 BUSI 和 CVC 两个数据集的主实验，U-KAN 均取得当前实验矩阵中的最佳 IoU 和 Dice；与 no-KAN 的对比说明 KANLinear 对分割性能具有实际贡献；与 U-Net 的对比说明 U-KAN 在精度上具有优势，但推理速度明显较慢。Attention-U-KAN 的实验结果未超过原始 U-KAN，提示简单注意力机制并非稳定改进方向。",
        "补充实验进一步增强了结论可信度。BUSI seed 6142 显示 U-KAN 相对 no-KAN 的优势仍然存在，但绝对指标受随机划分影响较大；CVC 继续训练实验没有超过原始 100 epoch best checkpoint，说明当前 U-KAN 结果与官方参考差距不能简单归因于训练不够。总体而言，本文不仅完成了算法复现，还通过消融实验、补充实验和负结果讨论体现了独立分析能力。最终项目形成了可运行代码、清晰 README、数据处理脚本、实验日志、图表资产和规范课程论文，满足课程方向（一）及多个加分项要求。",
    ]:
        add_para(doc, text)

    add_heading(doc, "参考文献", 1)
    refs = [
        "Li C., Liu X., Li W., et al. U-KAN Makes Strong Backbone for Medical Image Segmentation and Generation. arXiv preprint, 2024.",
        "Liu Z., Wang Y., Vaidya S., et al. KAN: Kolmogorov-Arnold Networks. arXiv:2404.19756, 2024.",
        "Ronneberger O., Fischer P., Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation. MICCAI, 2015.",
        "Woo S., Park J., Lee J.-Y., Kweon I. S. CBAM: Convolutional Block Attention Module. ECCV, 2018.",
        "Al-Dhabyani W., Gomaa M., Khaled H., Fahmy A. Dataset of Breast Ultrasound Images. Data in Brief, 2020.",
        "Bernal J., Sánchez F. J., Fernández-Esparrach G., et al. WM-DOVA maps for accurate polyp highlighting in colonoscopy. Computerized Medical Imaging and Graphics, 2015.",
        "CUHK-AIM-Group. U-KAN official GitHub repository. https://github.com/CUHK-AIM-Group/U-KAN",
        "Liu Z. pykan official GitHub repository. https://github.com/KindXiaoming/pykan",
        "A Review of KANs for Medical Image Analysis. 本地课程材料中的 KAN 医学图像分析综述资料。",
    ]
    for i, ref in enumerate(refs, 1):
        add_para(doc, f"[{i}] {ref}", first_line=False)

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def main() -> None:
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    figs = generate_figures()
    tables = build_table_data()
    stats = main_findings()
    build_markdown(figs, tables, stats)
    build_docx(figs, tables, stats)
    print(f"wrote_markdown={MD_OUT}")
    print(f"wrote_docx={DOCX_OUT}")
    print(f"wrote_figures={FIG_DIR}")


if __name__ == "__main__":
    main()
