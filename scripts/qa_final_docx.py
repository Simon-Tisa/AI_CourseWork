from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm, Pt


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit the generated course paper DOCX.")
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()

    path = args.docx.resolve()
    with ZipFile(path) as archive:
        zip_error = archive.testzip()
        media_count = sum(
            name.startswith("word/media/") for name in archive.namelist()
        )

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    reference_index = text.find("参考文献")
    body = text[:reference_index] if reference_index >= 0 else text
    captions = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style
        and paragraph.style.name == "Caption"
        and re.match(r"^图\s*\d+", paragraph.text.strip())
    ]
    figure_numbers = [
        int(re.match(r"^图\s*(\d+)", caption).group(1)) for caption in captions
    ]
    formula_numbers = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    match = re.fullmatch(r"（([456])-(\d+)）", paragraph.text.strip())
                    if match:
                        formula_numbers.append(
                            (int(match.group(1)), int(match.group(2)))
                        )
    chinese_chars = sum("\u4e00" <= character <= "\u9fff" for character in body)

    print(f"zip_test={zip_error}")
    print(f"bytes={path.stat().st_size}")
    print(f"sha256={hashlib.sha256(path.read_bytes()).hexdigest()}")
    print(f"inline_images={len(document.inline_shapes)}")
    print(f"media_files={media_count}")
    print(f"figure_captions={len(captions)}")
    print(f"figure_numbers={figure_numbers}")
    print(f"figures_sequential={figure_numbers == list(range(1, 25))}")
    print(f"formula_numbers={formula_numbers}")
    print(f"tables={len(document.tables)}")
    print(f"chinese_chars_before_references={chinese_chars}")

    assert zip_error is None
    # 24 numbered paper figures plus the Jinan University logo on the cover.
    assert len(document.inline_shapes) == 25
    assert media_count == 25
    assert figure_numbers == list(range(1, 25))
    assert formula_numbers == (
        [(4, number) for number in range(1, 10)]
        + [(5, number) for number in range(1, 5)]
        + [(6, number) for number in range(1, 9)]
    )
    assert chinese_chars >= 15000

    assert len(document.sections) == 3
    for section in document.sections:
        assert abs(section.page_width - Cm(21.0)) < Pt(0.5)
        assert abs(section.page_height - Cm(29.7)) < Pt(0.5)
        assert abs(section.top_margin - Cm(2.5)) < Pt(0.5)
        assert abs(section.bottom_margin - Cm(2.5)) < Pt(0.5)
        assert abs(section.left_margin - Cm(3.0)) < Pt(0.5)
        assert abs(section.right_margin - Cm(2.0)) < Pt(0.5)

    expected_styles = {
        "Normal": (12.0, False, 1.5),
        "Heading 1": (16.0, True, 1.0),
        "Heading 2": (15.0, True, 1.0),
        "Heading 3": (12.0, True, 1.0),
        "Caption": (10.5, False, 1.0),
    }
    for name, (size, bold, spacing) in expected_styles.items():
        style = document.styles[name]
        print(
            f"style.{name}=font:{style.font.name},size:{style.font.size.pt if style.font.size else None},"
            f"bold:{style.font.bold},spacing:{style.paragraph_format.line_spacing}"
        )
        assert style.font.name == "宋体"
        assert style.font.size is not None
        assert abs(style.font.size.pt - size) < 0.01
        assert bool(style.font.bold) is bold
        assert style.paragraph_format.line_spacing == spacing


if __name__ == "__main__":
    main()
